
import io
import os
import json
import hmac
import hashlib
import tempfile
import html
import re
import ssl
import socket
import time
import sqlite3
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, timedelta, datetime
from pathlib import Path

import pandas as pd
import requests
import streamlit as st
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from urllib3.poolmanager import PoolManager
from bs4 import BeautifulSoup


# =========================================================
# 기본 설정
# =========================================================
st.set_page_config(
    page_title="법인 Finder",
    page_icon="🏢",
    layout="wide",
)


DART_BASE = "https://opendart.fss.or.kr/api"
DART_ENG_BASE = "https://engopendart.fss.or.kr/engapi"
# 실행 위치와 관계없이 app.py와 같은 폴더에 DB를 저장합니다.
# 별도 경로가 필요하면 CORP_FINDER_DB 환경변수로 지정할 수 있습니다.
_db_override = os.getenv("CORP_FINDER_DB", "").strip()
DB_PATH = Path(_db_override).expanduser().resolve() if _db_override else (Path(__file__).resolve().parent / "corp_finder_cache.sqlite3")


# =========================================================
# 거래은행 표시 UI
# =========================================================
st.markdown(
    """
    <style>
    .bank-note {
        background: #f7f9fc;
        border: 1px solid #e1e6ef;
        border-left: 5px solid #ff4b4b;
        border-radius: 8px;
        padding: 12px 14px;
        margin: 8px 0 12px 0;
        line-height: 1.55;
        font-size: 0.97rem;
        color: #1f2937;
        white-space: pre-wrap;
    }
    .bank-meta {
        display: flex;
        flex-wrap: wrap;
        gap: 7px;
        margin: 4px 0 10px 0;
    }
    .bank-chip {
        display: inline-block;
        padding: 4px 9px;
        border-radius: 999px;
        background: #eef2f7;
        border: 1px solid #dde3ec;
        font-size: 0.86rem;
        font-weight: 650;
        color: #273142;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


def _evidence_items(bank_row):
    raw = str(bank_row.get("근거문구") or "").strip()
    if not raw:
        return []
    parts = [re.sub(r"\\s+", " ", x).strip() for x in re.split(r"\\n\\s*\\n", raw) if x.strip()]
    seen, out = set(), []
    for x in parts:
        if x not in seen:
            seen.add(x)
            out.append(x)
    return out


def render_bank_analysis(banks, title="🏦 감사보고서에서 확인된 거래 금융기관"):
    """거래은행 요약표와 은행별 감사보고서 근거를 한 화면에 표시한다."""
    if not banks:
        return

    st.markdown(f"### {title}")

    summary_rows = []
    for b in banks:
        evidences = _evidence_items(b)
        representative = evidences[0] if evidences else "-"
        if len(representative) > 180:
            representative = representative[:177] + "..."
        summary_rows.append({
            "은행": b.get("은행") or "-",
            "거래유형": b.get("거래유형") or "-",
            "확인강도": b.get("확인강도") or "-",
            "금액표기": b.get("금액표기") or "-",
            "근거건수": b.get("근거건수") or len(evidences),
            "대표 근거문구": representative,
        })

    summary_df = pd.DataFrame(summary_rows)
    table_height = min(90 + max(len(summary_df), 1) * 58, 620)
    st.dataframe(
        summary_df,
        use_container_width=True,
        hide_index=True,
        height=table_height,
        row_height=52,
        column_config={
            "은행": st.column_config.TextColumn("은행", width="medium"),
            "거래유형": st.column_config.TextColumn("거래유형", width="large"),
            "확인강도": st.column_config.TextColumn("확인강도", width="small"),
            "금액표기": st.column_config.TextColumn("금액표기", width="medium"),
            "근거건수": st.column_config.NumberColumn("근거건수", width="small", format="%d건"),
            "대표 근거문구": st.column_config.TextColumn("감사보고서 대표 근거", width="large"),
        },
    )

    st.caption(
        "※ 거래유형은 차입/대출·담보/질권·예금·보증/약정 등의 문맥을 함께 분석한 결과입니다. "
        "은행명이 기재되었다는 사실만으로 주거래은행이라고 단정하지 않습니다."
    )

    st.markdown("#### 📑 은행별 거래유형 · 감사보고서 근거")
    st.caption("아래 근거문구는 접어두지 않고 바로 표시됩니다. 은행별로 실제 보고서 문맥을 한 화면에서 비교할 수 있습니다.")

    for b in banks:
        evidences = _evidence_items(b)
        with st.container(border=True):
            head_left, head_right = st.columns([5, 1])
            with head_left:
                st.markdown(f"### 🏦 {b.get('은행') or '-'}")
                chips = []
                for value in [b.get("거래유형"), f"확인강도 {b.get('확인강도') or '-'}"]:
                    if value:
                        chips.append(f'<span class="bank-chip">{html.escape(str(value))}</span>')
                if b.get("금액표기"):
                    chips.append(f'<span class="bank-chip">금액 {html.escape(str(b.get("금액표기")))}</span>')
                st.markdown('<div class="bank-meta">' + ''.join(chips) + '</div>', unsafe_allow_html=True)
            with head_right:
                st.metric("근거", f"{b.get('근거건수') or len(evidences)}건")

            if evidences:
                for idx, evidence in enumerate(evidences, start=1):
                    safe = html.escape(evidence)
                    st.markdown(
                        f'<div class="bank-note"><b>근거 {idx}</b><br>{safe}</div>',
                        unsafe_allow_html=True,
                    )
            else:
                st.info("표시할 감사보고서 근거문구가 없습니다.")


# =========================================================
# OpenDART HTTPS 연결 (국문 -> 공식 영문 경로, 장애 시 단기 차단)
# =========================================================
# 접속 제한 3~4초가 정상적인 느린 연결도 실패 처리하던 문제를 개선한다.
# 두 서버 모두 접속 불가능할 경우 같은 세션에서 반복 접속을 잠시 방지한다.
# DART 측 IP/네트워크 차단은 애플리케이션만 수정해서 해제할 수 없다.
_DART_CONNECT_TIMEOUT = 8
_DART_DIAG_CONNECT_TIMEOUT = 5
_DART_NETWORK_COOLDOWN = 90


class DARTSSLAdapter(HTTPAdapter):
    def init_poolmanager(self, connections, maxsize, block=False, **pool_kwargs):
        # 레거시 DH_KEY_TOO_SMALL을 겪는 PC에 한해 사용하는 호환 어댑터.
        ctx = ssl.create_default_context()
        ctx.set_ciphers("DEFAULT:@SECLEVEL=1")
        pool_kwargs["ssl_context"] = ctx
        self.poolmanager = PoolManager(
            num_pools=connections, maxsize=maxsize, block=block, **pool_kwargs
        )


def _retry_policy():
    # 수동으로 국문/영문 우회를 관리하므로 requests의 숨은 재시도는 사용하지 않는다.
    return Retry(total=0, connect=0, read=0, redirect=0, backoff_factor=0)


@st.cache_resource
def get_dart_default_session():
    session = requests.Session()
    session.mount("https://", HTTPAdapter(max_retries=_retry_policy()))
    session.headers.update({
        "User-Agent": "Mozilla/5.0 Corporate-Finder/2.9",
        "Accept": "*/*",
        "Connection": "keep-alive",
    })
    return session


@st.cache_resource
def get_dart_primary_session():
    return get_dart_default_session()


@st.cache_resource
def get_dart_legacy_session():
    session = requests.Session()
    session.mount(
        "https://opendart.fss.or.kr/",
        DARTSSLAdapter(max_retries=_retry_policy()),
    )
    session.headers.update({
        "User-Agent": "Mozilla/5.0 Corporate-Finder/2.9",
        "Accept": "*/*",
        "Connection": "keep-alive",
    })
    return session


def _to_eng_dart_url(url: str) -> str:
    """국문 OpenDART API URL을 공식 영문 OpenDART API URL로 변환."""
    if isinstance(url, str) and url.startswith("https://opendart.fss.or.kr/api/"):
        return url.replace(
            "https://opendart.fss.or.kr/api/",
            "https://engopendart.fss.or.kr/engapi/", 1,
        )
    return url


def _dart_safe_error(exc: Exception) -> str:
    """인증키/요청 URL/내부 예외의 민감한 텍스트는 절대 노출하지 않는다."""
    if isinstance(exc, requests.exceptions.ConnectTimeout):
        return "연결 시간 초과"
    if isinstance(exc, requests.exceptions.ReadTimeout):
        return "서버 응답 시간 초과"
    if isinstance(exc, requests.exceptions.SSLError):
        return "TLS/SSL 인증 또는 협상 실패"
    if isinstance(exc, requests.exceptions.ProxyError):
        return "중간 프록시 연결 실패"
    if isinstance(exc, requests.exceptions.ConnectionError):
        # requests가 ConnectionError 안에 DNSError 등을 감싸는 경우가 있음.
        detail = str(exc).lower()
        if any(token in detail for token in ("nameresolutionerror", "name or service not known", "temporary failure in name resolution")):
            return "DNS 이름 조회 실패"
        return "네트워크 연결 실패"
    if isinstance(exc, requests.exceptions.HTTPError):
        response = getattr(exc, "response", None)
        return f"HTTP {response.status_code}" if response is not None else "HTTP 응답 오류"
    if isinstance(exc, requests.exceptions.RequestException):
        return type(exc).__name__
    return type(exc).__name__


def _check_dart_http(response):
    if not (200 <= response.status_code < 300):
        raise RuntimeError(f"DART HTTP 응답 오류 ({response.status_code})")


def _dart_cooldown_remaining():
    """실패한 서버에 매 화면 렌더링마다 재접속하는 현상 방지."""
    return max(0, int(float(st.session_state.get("_dart_cooldown_until", 0)) - time.monotonic()))


def _dart_reset_cooldown():
    st.session_state.pop("_dart_cooldown_until", None)
    st.session_state.pop("_dart_last_network_error", None)


def _dart_mark_network_failure(message):
    st.session_state["_dart_cooldown_until"] = time.monotonic() + _DART_NETWORK_COOLDOWN
    st.session_state["_dart_last_network_error"] = message


def dart_get(url, **kwargs):
    """국문 API 요청; 접속 실패 시 영문 API로 우회. 무의미한 반복 대기 방지."""
    remaining = _dart_cooldown_remaining()
    if remaining:
        raise RuntimeError(
            f"DART 국문·영문 연결 실패 후 재접속 대기 중({remaining}초). "
            "'DART 연결 진단'으로 즉시 다시 확인할 수 있습니다."
        )

    original_timeout = kwargs.get("timeout", 30)
    if isinstance(original_timeout, (tuple, list)):
        read_timeout = float(original_timeout[1])
    else:
        read_timeout = float(original_timeout)
    request_kwargs = dict(kwargs)
    request_kwargs["timeout"] = (_DART_CONNECT_TIMEOUT, read_timeout)

    primary_error = None
    try:
        response = get_dart_primary_session().get(url, **request_kwargs)
        if response.status_code < 500:
            _dart_reset_cooldown()
            return response
        primary_error = f"HTTP {response.status_code}"
    except requests.exceptions.SSLError as exc:
        if "DH_KEY_TOO_SMALL" in str(exc).upper() or "DH KEY TOO SMALL" in str(exc).upper():
            try:
                response = get_dart_legacy_session().get(url, **request_kwargs)
                if response.status_code < 500:
                    _dart_reset_cooldown()
                    return response
                primary_error = f"HTTP {response.status_code}"
            except requests.exceptions.RequestException as retry_exc:
                primary_error = _dart_safe_error(retry_exc)
        else:
            primary_error = _dart_safe_error(exc)
    except requests.exceptions.RequestException as exc:
        primary_error = _dart_safe_error(exc)

    alt = _to_eng_dart_url(url)
    if alt == url:
        _dart_mark_network_failure(primary_error or "국문 서버 접속 실패")
        raise RuntimeError("DART 국문 서버 연결 실패: " + (primary_error or "알 수 없는 오류"))

    secondary_error = None
    try:
        response = get_dart_default_session().get(alt, **request_kwargs)
        if response.status_code < 500:
            _dart_reset_cooldown()
            return response
        secondary_error = f"HTTP {response.status_code}"
    except requests.exceptions.RequestException as exc:
        secondary_error = _dart_safe_error(exc)

    error_message = (
        f"국문: {primary_error or '접속 실패'} / "
        f"영문: {secondary_error or '접속 실패'}"
    )
    _dart_mark_network_failure(error_message)
    raise RuntimeError(
        f"DART 국문·영문 서버 연결 실패 ({error_message}). "
        "네트워크·IP 제한 또는 서버 응답 지연 여부를 연결 진단에서 확인해 주세요."
    )


def check_dart_connection(api_key: str, corp_code: str):
    """DNS/TCP/TLS/API를 나누어 테스트한다. 인증키/실제 요청 URL은 표시하지 않는다.

    TCP 테스트는 직접 연결이며, HTTP 테스트는 requests의 시스템 프록시 설정을 이용할 수 있다.
    단계별 실패가 반드시 DART의 장애를 뜻하지는 않는다.
    """
    endpoints = [
        ("국문", "opendart.fss.or.kr", f"{DART_BASE}/list.json"),
        ("영문", "engopendart.fss.or.kr", f"{DART_ENG_BASE}/list.json"),
    ]
    results = []
    params = {
        "crtfc_key": api_key,
        "corp_code": corp_code,
        "bgn_de": (date.today() - timedelta(days=365)).strftime("%Y%m%d"),
        "end_de": date.today().strftime("%Y%m%d"),
        "page_count": 1,
    }
    for label, host, endpoint in endpoints:
        began = time.perf_counter()
        try:
            socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
            dns_ok = True
            dns_state = "정상"
        except (OSError, ValueError):
            dns_ok = False
            dns_state = "이름 조회 실패"
        results.append({"서버": label, "단계": "DNS", "결과": dns_state,
                        "소요시간(초)": round(time.perf_counter() - began, 1)})

        raw_sock = None
        began = time.perf_counter()
        tcp_ok = False
        if dns_ok:
            try:
                raw_sock = socket.create_connection((host, 443), timeout=_DART_DIAG_CONNECT_TIMEOUT)
                tcp_ok = True
                tcp_state = "정상"
            except (OSError, ValueError):
                tcp_state = "접속 실패/시간 초과"
        else:
            tcp_state = "DNS 실패로 생략"
        results.append({"서버": label, "단계": "TCP 443", "결과": tcp_state,
                        "소요시간(초)": round(time.perf_counter() - began, 1)})

        began = time.perf_counter()
        if tcp_ok:
            try:
                with ssl.create_default_context().wrap_socket(raw_sock, server_hostname=host) as tls_sock:
                    tls_sock.settimeout(_DART_DIAG_CONNECT_TIMEOUT)
                    tls_state = "정상"
            except (ssl.SSLError, OSError, ValueError):
                tls_state = "TLS 협상 실패"
            finally:
                if raw_sock:
                    raw_sock.close()
        else:
            tls_state = "TCP 실패로 생략"
        results.append({"서버": label, "단계": "TLS", "결과": tls_state,
                        "소요시간(초)": round(time.perf_counter() - began, 1)})

        began = time.perf_counter()
        try:
            r = get_dart_default_session().get(
                endpoint, params=params, timeout=(_DART_DIAG_CONNECT_TIMEOUT, 12)
            )
            if r.ok:
                # 실제 HTTP 응답이 있으면 네트워크 자체는 도달 가능한 상태.
                # DART 012(IP 제한) 등은 다음 요청에서 구체적인 원인으로 표시한다.
                _dart_reset_cooldown()
                try:
                    data = r.json()
                    code = str(data.get("status", "미확인"))
                    state = f"HTTP {r.status_code} / DART {code}"
                    if code in ("000", "013"):
                        _dart_reset_cooldown()
                    elif code == "012":
                        state += " (접속 IP 제한)"
                    elif code == "020":
                        state += " (요청 한도)"
                    elif code in ("010", "011", "901"):
                        state += " (인증키 상태 확인)"
                except (ValueError, AttributeError):
                    state = f"HTTP {r.status_code} / JSON 해석 불가"
            else:
                state = f"HTTP {r.status_code}"
        except requests.exceptions.RequestException as exc:
            state = _dart_safe_error(exc)
        results.append({"서버": label, "단계": "공시목록 API", "결과": state,
                        "소요시간(초)": round(time.perf_counter() - began, 1)})
    return results


BANK_PATTERNS = {
    "하나은행": [
        r"하나은행", r"KEB\s*하나은행", r"KEB하나", r"㈜하나은행", r"\(주\)하나은행",
        r"Hana\s*Bank", r"KEB\s*Hana\s*Bank"
    ],
    "KB국민은행": [
        r"KB\s*국민은행", r"국민은행", r"㈜국민은행", r"\(주\)국민은행",
        r"KB\s*Kookmin\s*Bank", r"Kookmin\s*Bank"
    ],
    "신한은행": [r"신한은행", r"㈜신한은행", r"\(주\)신한은행", r"Shinhan\s*Bank"],
    "우리은행": [r"우리은행", r"㈜우리은행", r"\(주\)우리은행", r"Woori\s*Bank"],
    "NH농협은행": [
        r"NH\s*농협은행", r"농협은행", r"농협중앙회", r"NH농협",
        r"NH\s*NongHyup\s*Bank", r"NongHyup\s*Bank", r"National\s*Agricultural\s*Cooperative"
    ],
    "IBK기업은행": [
        r"IBK\s*기업은행", r"기업은행", r"중소기업은행",
        r"Industrial\s*Bank\s*of\s*Korea", r"IBK"
    ],
    "KDB산업은행": [
        r"KDB\s*산업은행", r"산업은행", r"한국산업은행",
        r"Korea\s*Development\s*Bank", r"KDB"
    ],
    "SC제일은행": [
        r"SC\s*제일은행", r"제일은행", r"스탠다드차타드은행", r"한국스탠다드차타드은행",
        r"Standard\s*Chartered\s*Bank\s*Korea", r"Standard\s*Chartered"
    ],
    "한국씨티은행": [r"한국씨티은행", r"씨티은행", r"Citibank\s*Korea", r"Citibank"],
    "iM뱅크(구 대구은행)": [r"iM뱅크", r"아이엠뱅크", r"대구은행", r"iM\s*Bank", r"Daegu\s*Bank"],
    "BNK부산은행": [r"BNK\s*부산은행", r"부산은행", r"Busan\s*Bank"],
    "BNK경남은행": [r"BNK\s*경남은행", r"경남은행", r"Kyongnam\s*Bank", r"Gyeongnam\s*Bank"],
    "광주은행": [r"광주은행", r"Kwangju\s*Bank", r"Gwangju\s*Bank"],
    "전북은행": [r"전북은행", r"Jeonbuk\s*Bank"],
    "제주은행": [r"제주은행", r"Jeju\s*Bank"],
    "Sh수협은행": [r"Sh\s*수협은행", r"수협은행", r"Suhyup\s*Bank", r"National\s*Federation\s*of\s*Fisheries"],
}

CATEGORY_PATTERNS = {
    "차입/대출": [
        r"차입금", r"단기차입", r"장기차입", r"대출", r"시설자금",
        r"운전자금", r"일반자금", r"한도대출",
        r"borrowings?", r"loan", r"credit\s*line", r"working\s*capital", r"facility"
    ],
    "담보/질권": [
        r"담보", r"근저당", r"질권", r"담보제공", r"담보설정",
        r"collateral", r"pledge", r"mortgage", r"security\s*interest"
    ],
    "예금": [
        r"예금", r"정기예금", r"사용제한", r"금융상품", r"예치금",
        r"deposit", r"restricted\s*deposit", r"financial\s*instrument"
    ],
    "보증/약정": [
        r"지급보증", r"보증", r"약정", r"한도약정", r"신용장", r"외화",
        r"guarantee", r"commitment", r"letter\s*of\s*credit", r"credit\s*facility"
    ],
}

RELATED_WORDS = [
    "은행", "금융기관", "차입금", "단기차입금", "장기차입금",
    "담보", "질권", "예금", "정기예금", "사용제한예금",
    "대출", "지급보증", "한도약정", "약정",
    "bank", "financial institution", "borrowing", "loan", "collateral",
    "pledge", "deposit", "guarantee", "commitment", "credit facility"
]


# =========================================================
# 유틸리티
# =========================================================
def get_secret(name, default=""):
    # Streamlit Cloud: st.secrets / 로컬: st.secrets 또는 환경변수 모두 지원
    try:
        value = st.secrets.get(name, None)
        if value not in (None, ""):
            return value
    except Exception:
        pass
    return os.getenv(name, default)


def save_external_company_db(df: pd.DataFrame, mapped: dict):
    """외부 XLSX/CSV의 정규화된 법인목록까지 SQLite에 영구 저장합니다."""
    if df is None or df.empty:
        raise ValueError('저장할 기업목록 데이터가 없습니다.')
    with get_db_connection() as conn:
        df.fillna('').astype(str).to_sql('external_company', conn, if_exists='replace', index=False)
        conn.execute('CREATE INDEX IF NOT EXISTS idx_external_corp_name ON external_company(norm_name)')
        conn.execute("INSERT OR REPLACE INTO db_meta(meta_key,meta_value) VALUES('external_company_meta',?)",
                     (json.dumps(mapped, ensure_ascii=False, default=str),))
        conn.execute("INSERT OR REPLACE INTO db_meta(meta_key,meta_value) VALUES('external_company_updated_at',?)",
                     (datetime.now().strftime('%Y-%m-%d %H:%M:%S'),))
    _db_cloud_push()


def load_external_company_db():
    with get_db_connection() as conn:
        found=conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='external_company'").fetchone()
        if not found:
            return pd.DataFrame(), {}
        df=pd.read_sql_query('SELECT * FROM external_company',conn).fillna('')
    try:
        meta=json.loads(get_db_meta('external_company_meta','{}'))
    except (ValueError, TypeError):
        meta={}
    return df, meta


def normalize_company_name(name: str) -> str:
    if not name:
        return ""
    x = re.sub(r"\s+", "", str(name))
    x = x.replace("주식회사", "").replace("(주)", "").replace("㈜", "")
    return x.lower()


def request_json(url, params, timeout=30):
    r = dart_get(url, params=params, timeout=timeout)
    _check_dart_http(r)
    data = r.json()
    status = str(data.get("status", "000"))
    # 013은 조회된 데이터가 없다는 정상적인 빈 결과 코드
    if status == "013":
        return data
    if status not in ("000", ""):
        raise RuntimeError(data.get("message", f"DART 오류 코드: {status}"))
    return data


def decode_bytes(raw: bytes) -> str:
    for enc in ("utf-8", "cp949", "euc-kr"):
        try:
            return raw.decode(enc)
        except Exception:
            pass
    return raw.decode("utf-8", errors="ignore")


def clean_text(raw_text: str) -> str:
    soup = BeautifulSoup(raw_text, "html.parser")
    text = soup.get_text("\n")
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def money_tokens(text: str):
    # 감사보고서에서 흔한 금액 표기를 넓게 포착
    patt = r"(?<!\d)(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\s*(?:원|천원|백만원|억원)"
    found = re.findall(patt, text)
    result = []
    for x in found:
        x = re.sub(r"\s+", "", x)
        if x not in result:
            result.append(x)
    return result[:5]



# =========================================================
# 지역검색용 로컬 DB / 외부 법인목록
# =========================================================
REGION_PROVINCES = [
    "전체", "서울특별시", "부산광역시", "대구광역시", "인천광역시",
    "광주광역시", "대전광역시", "울산광역시", "세종특별자치시",
    "경기도", "강원특별자치도", "충청북도", "충청남도", "전북특별자치도",
    "전라남도", "경상북도", "경상남도", "제주특별자치도",
]

CHEONGJU_DISTRICTS = ["전체", "상당구", "서원구", "흥덕구", "청원구"]

EXTERNAL_COL_CANDIDATES = {
    "법인명": ["법인명", "회사명", "업체명", "기업명", "상호명", "상호", "업소명"],
    "대표자": ["대표자", "대표자명", "대표"],
    "주소": [
        "본점주소", "본사주소", "공장대표주소(도로명)", "공장대표주소", "사업장 주소",
        "사업장주소", "영업소 주소(도로명)", "영업소주소", "소재지", "주소", "도로명주소"
    ],
    "지역": ["시군", "시군구", "지역", "소재 시군", "소재시군"],
    "업종": ["업종", "업종명", "산업분류", "산업분류명", "KSIC", "주생산품", "제품"],
    "전화번호": ["전화번호", "전화", "소재지전화", "사업장 번호", "사업장번호"],
    "사업자등록번호": ["사업자등록번호", "사업자번호"],
    "법인등록번호": ["법인등록번호", "법인번호"],
    "종업원수": ["상시근로자수", "종업원수", "직원수", "근로자수", "임직원수"],
}


class _AutoClosingSQLiteConnection(sqlite3.Connection):
    """sqlite3.Connection.__exit__ commits/rolls back but does NOT close.

    In Windows leaving these handles open prevents deleting temporary SQLite
    snapshots (WinError 32) and can keep the live database locked as well.
    """

    def __exit__(self, exc_type, exc_value, traceback):
        try:
            return super().__exit__(exc_type, exc_value, traceback)
        finally:
            self.close()


def get_db_connection():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(
        DB_PATH, check_same_thread=False, factory=_AutoClosingSQLiteConnection
    )
    conn.row_factory = sqlite3.Row
    return conn


def init_local_db():
    with get_db_connection() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS corp_master (
                corp_code TEXT PRIMARY KEY,
                corp_name TEXT NOT NULL,
                stock_code TEXT,
                modify_date TEXT,
                norm_name TEXT
            );

            CREATE TABLE IF NOT EXISTS company_profile (
                corp_code TEXT PRIMARY KEY,
                corp_name TEXT,
                stock_code TEXT,
                modify_date TEXT,
                ceo_nm TEXT,
                corp_cls TEXT,
                jurir_no TEXT,
                bizr_no TEXT,
                adres TEXT,
                hm_url TEXT,
                ir_url TEXT,
                phn_no TEXT,
                fax_no TEXT,
                induty_code TEXT,
                fetched_at TEXT,
                fetch_error TEXT DEFAULT ''
            );

            CREATE INDEX IF NOT EXISTS idx_profile_adres ON company_profile(adres);
            CREATE INDEX IF NOT EXISTS idx_profile_name ON company_profile(corp_name);
            CREATE INDEX IF NOT EXISTS idx_master_norm_name ON corp_master(norm_name);

            CREATE TABLE IF NOT EXISTS db_meta (
                meta_key TEXT PRIMARY KEY,
                meta_value TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS company_analysis (
                corp_code TEXT PRIMARY KEY,
                result_json TEXT NOT NULL,
                analyzed_at TEXT NOT NULL
            );
            """
        )
    # DB 스냅샷/기존 Cloud DB는 새 퇴직연금 테이블이 없을 수 있습니다.
    # 앱 시작/복원 시 CREATE TABLE IF NOT EXISTS 로 안전하게 스키마만 추가합니다.
    # 기존 법인 데이터, 분석 결과, 영업 메모는 삭제하거나 초기화하지 않습니다.
    _pension_target_init()
    _pension_employee_init()


def set_db_meta(key: str, value: str):
    with get_db_connection() as conn:
        conn.execute(
            "INSERT INTO db_meta(meta_key, meta_value) VALUES (?, ?) "
            "ON CONFLICT(meta_key) DO UPDATE SET meta_value=excluded.meta_value",
            (key, value),
        )


def get_db_meta(key: str, default: str = "") -> str:
    with get_db_connection() as conn:
        row = conn.execute("SELECT meta_value FROM db_meta WHERE meta_key=?", (key,)).fetchone()
    return str(row[0]) if row else default


def load_saved_corp_master() -> pd.DataFrame:
    """DART에 접속하지 않고 마지막으로 저장한 법인 고유번호 목록을 읽습니다."""
    with get_db_connection() as conn:
        df = pd.read_sql_query(
            "SELECT corp_code, corp_name, stock_code, modify_date, norm_name "
            "FROM corp_master ORDER BY corp_name", conn,
        )
    if df.empty:
        return pd.DataFrame(columns=["corp_code", "corp_name", "stock_code", "modify_date", "norm_name"])
    for col in ("corp_code", "corp_name", "stock_code", "modify_date", "norm_name"):
        df[col] = df[col].fillna("").astype(str)
    missing = df["norm_name"].eq("")
    if missing.any():
        df.loc[missing, "norm_name"] = df.loc[missing, "corp_name"].map(normalize_company_name)
    return df


def get_saved_company_profile(corp_code: str) -> dict | None:
    with get_db_connection() as conn:
        row = conn.execute(
            "SELECT * FROM company_profile WHERE corp_code=?", (str(corp_code),)
        ).fetchone()
    if row is None or (row["fetch_error"] and not row["adres"]):
        return None
    return dict(row)


def save_company_analysis(result: dict):
    code = str((result.get("company") or {}).get("DART고유번호", "")).strip()
    if not code:
        return
    now = datetime.now().isoformat(timespec="seconds")
    payload = json.dumps(result, ensure_ascii=False, default=str)
    with get_db_connection() as conn:
        conn.execute(
            "INSERT INTO company_analysis(corp_code, result_json, analyzed_at) VALUES(?,?,?) "
            "ON CONFLICT(corp_code) DO UPDATE SET "
            "result_json=excluded.result_json, analyzed_at=excluded.analyzed_at",
            (code, payload, now),
        )
    _db_cloud_push()


def get_saved_company_analysis(corp_code: str):
    with get_db_connection() as conn:
        row = conn.execute(
            "SELECT result_json, analyzed_at FROM company_analysis WHERE corp_code=?",
            (str(corp_code),),
        ).fetchone()
    if not row:
        return None, ""
    return json.loads(row["result_json"]), row["analyzed_at"]


def export_full_db_bytes() -> bytes:
    """백업 시점의 전체 SQLite DB 스냅샷(법인/주소/분석 결과)을 생성합니다."""
    with tempfile.NamedTemporaryFile(suffix=".sqlite3", delete=False) as f:
        tmp = Path(f.name)
    src = dst = None
    try:
        src = get_db_connection()
        dst = sqlite3.connect(tmp)
        src.backup(dst)
        dst.close(); dst = None
        src.close(); src = None
        return tmp.read_bytes()
    finally:
        if src is not None: src.close()
        if dst is not None: dst.close()
        tmp.unlink(missing_ok=True)


def restore_full_db_bytes(raw: bytes):
    """스키마/무결성 검사 후 백업 DB를 복원합니다."""
    if not raw or not raw.startswith(b"SQLite format 3\x00"):
        raise ValueError("올바른 SQLite DB 파일이 아닙니다.")
    with tempfile.NamedTemporaryFile(suffix=".sqlite3", delete=False) as f:
        tmp = Path(f.name)
        f.write(raw)
    src = dst = None
    try:
        src = sqlite3.connect(tmp)
        integrity = src.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise ValueError(f"DB 무결성 검증 실패: {integrity}")
        tbls = {r[0] for r in src.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if not {"corp_master", "company_profile"}.issubset(tbls):
            raise ValueError("Corporate Finder의 법인/주소 DB가 아닙니다.")
        dst = get_db_connection()
        src.backup(dst)
        dst.close(); dst = None
        src.close(); src = None
        init_local_db()  # 이전 버전 백업도 최신 테이블로 자동 마이그레이션
    finally:
        if src is not None: src.close()
        if dst is not None: dst.close()
        tmp.unlink(missing_ok=True)


# =========================================================
# HANAV CLOUD SQLITE SYNC — PRIVATE SUPABASE STORAGE
# Last good snapshot is kept remotely, and restored on a new server start.
# SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY, SUPABASE_BUCKET are server secrets.
# Never expose the service-role key in the repository or the UI.
# =========================================================
import hashlib as _db_hashlib
import gzip as _db_gzip
import tempfile as _db_tempfile
from contextlib import closing as _db_closing
from urllib.parse import quote as _db_quote

_DB_CLOUD_LEGACY_OBJECT = 'corporate/corp_finder_cache.sqlite3'
_DB_CLOUD_OBJECT = 'corporate/corp_finder_cache_v2.sqlite3.gz'
_DB_CLOUD_CONFIG_ERROR = ''


def _db_cloud_config():
    url = str(get_secret('SUPABASE_URL') or '').strip().rstrip('/')
    key = str(get_secret('SUPABASE_SERVICE_ROLE_KEY') or '').strip()
    bucket = str(get_secret('SUPABASE_BUCKET') or 'hanav-finder-db').strip()
    if not url or not key:
        return None
    if not url.startswith('https://'):
        raise RuntimeError('SUPABASE_URL은 https:// 주소여야 합니다.')
    if not bucket or '/' in bucket or '..' in bucket:
        raise RuntimeError('SUPABASE_BUCKET 설정을 확인해 주세요.')
    prefix = url + '/storage/v1/object/'
    suffix = _db_quote(bucket, safe='') + '/' + _db_quote(_DB_CLOUD_OBJECT, safe='/')
    legacy_suffix = _db_quote(bucket, safe='') + '/' + _db_quote(_DB_CLOUD_LEGACY_OBJECT, safe='/')
    return {'read_url': prefix + 'authenticated/' + suffix,
             'write_url': prefix + suffix,
             'legacy_read_url': prefix + 'authenticated/' + legacy_suffix,
             'headers': {'apikey': key, 'Authorization': 'Bearer ' + key},
             'bucket': bucket}



def _db_download_snapshot(cfg, *, allow_missing=False):
    """Read compressed v2 backup; fall back to *read-only* legacy .sqlite3.

    New backups never overwrite the old legacy object. Both versions are
    validated by their callers before attaching to the live database.
    """
    for url, compressed in ((cfg['read_url'], True),
                            (cfg['legacy_read_url'], False)):
        response = requests.get(url, headers=cfg['headers'], timeout=(6, 45))
        if response.status_code == 404:
            continue
        response.raise_for_status()
        if compressed:
            try:
                raw = _db_gzip.decompress(response.content)
            except (_db_gzip.BadGzipFile, EOFError, OSError) as exc:
                raise RuntimeError('압축된 Supabase 백업 파일이 손상되었습니다.') from exc
        else:
            raw = response.content
        if not raw.startswith(b'SQLite format 3\x00'):
            raise RuntimeError('Supabase 백업 파일 형식이 올바르지 않습니다.')
        return raw
    if allow_missing:
        return None
    raise RuntimeError('Supabase에 저장된 법인 DB 백업이 없습니다.')


def _db_valid_snapshot(raw):
    if not raw or not raw.startswith(b'SQLite format 3\x00'):
        raise ValueError('백업 파일이 SQLite 형식이 아닙니다.')
    with _db_tempfile.NamedTemporaryFile(suffix='.sqlite3', delete=False) as f:
        f.write(raw)
        tmp = Path(f.name)
    try:
        with _db_closing(sqlite3.connect(str(tmp))) as conn:
            if conn.execute('PRAGMA integrity_check').fetchone()[0] != 'ok':
                raise ValueError('DB 무결성 검사에 실패했습니다.')
            tables = {x[0] for x in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )}
            required = {'corp_master', 'company_profile'}
            if not required.issubset(tables):
                raise ValueError('다른 종류의 DB 파일이 저장되어 있습니다.')
    finally:
        tmp.unlink(missing_ok=True)


def _db_sqlite_snapshot():
    path = DB_PATH
    if not path.exists():
        raise FileNotFoundError('저장된 로컬 DB가 없습니다.')
    with _db_tempfile.NamedTemporaryFile(suffix='.sqlite3', delete=False) as f:
        tmp = Path(f.name)
    try:
        with _db_closing(sqlite3.connect(str(path), timeout=30)) as src, \
                _db_closing(sqlite3.connect(str(tmp))) as dst:
            src.backup(dst)  # WAL에 아직 남아 있는 데이터까지 포함
        # src/dst both CLOSED before reading / unlinking the temp file (Windows).
        raw = tmp.read_bytes()
        _db_valid_snapshot(raw)
        return raw
    finally:
        tmp.unlink(missing_ok=True)


def _db_install_snapshot(raw):
    # restore_full_db_bytes already verifies and stages its own temporary DB.
    # Avoid creating an extra identical temporary file that must be deleted.
    restore_full_db_bytes(raw)



def _db_cloud_push(show_feedback=False):
    """Validate snapshot, gzip it, and upload as a *new* Storage object.

    Keeps the previous .sqlite3 backup untouched in case first migration
    fails. Does not report success until Storage returns an HTTP success.
    """
    cfg = _db_cloud_config()
    if cfg is None:
        if show_feedback:
            st.warning('Supabase 저장 설정이 없습니다.')
        return False
    raw_size = 0
    zip_size = 0
    try:
        # Retain results from another PC/session when the existing backup is readable.
        try:
            _db_cloud_merge_analysis_from_remote()
        except Exception:
            pass
        raw = _db_sqlite_snapshot()  # includes a full integrity check
        raw_size = len(raw)
        packed = _db_gzip.compress(raw, compresslevel=6, mtime=0)
        zip_size = len(packed)
        headers = dict(cfg['headers'])
        headers.update({'Content-Type': 'application/gzip',
                        'x-upsert': 'true', 'cache-control': 'no-cache'})
        response = requests.post(cfg['write_url'], headers=headers, data=packed,
                                 timeout=(15, 180))
        response.raise_for_status()
        st.session_state['_db_cloud_saved_at'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        st.session_state.pop('_db_cloud_error', None)
        if show_feedback:
            st.success(f'☁️ DB 압축 백업 완료: {raw_size/1048576:.1f}MB → '
                       f'{zip_size/1048576:.1f}MB (Supabase 저장 성공)')
        return True
    except Exception as exc:
        # Exception text never embeds API keys in this path; display only safe type.
        if isinstance(exc, requests.exceptions.Timeout) or 'timed out' in str(exc).lower():
            reason = '클라우드 전송 시간 초과 (Supabase 연결 또는 네트워크 문제)'
        elif isinstance(exc, requests.exceptions.HTTPError):
            status = getattr(getattr(exc, 'response', None), 'status_code', '알 수 없음')
            reason = f'Supabase HTTP {status}'
        elif isinstance(exc, requests.exceptions.ConnectionError):
            reason = 'Supabase 네트워크 연결 오류'
        else:
            reason = f'{type(exc).__name__}: {exc}'
        st.session_state['_db_cloud_error'] = reason
        if show_feedback:
            st.error(f'클라우드 DB 백업 실패: {reason}. 로컬 DB는 유지됩니다.')
            if raw_size:
                st.caption(f'전송 데이터: 원본 {raw_size/1048576:.1f}MB → '
                           f'압축 {zip_size/1048576:.1f}MB')
        return False


def _db_cloud_merge_analysis_from_remote(only_corp: str | None = None) -> int:
    """Supabase 전체 DB 중 *분석 결과*만 현재 DB에 안전하게 병합한다.

    서버 실행시 1회 복원과 달리 실행 중에도 PC가 올린 최신 분석을 가져온다.
    기존 로컬 분석은 더 최신이거나 같으면 그대로 둔다. DB 전체를 덮어쓰지 않는다.
    서버의 다른 작업에 영향을 주지 않도록 쓰기/되올리기는 하지 않는다.
    """
    cfg = _db_cloud_config()
    if cfg is None:
        raise RuntimeError('Supabase 비공개 DB 연결 설정이 없습니다.')
    if only_corp is not None and not re.fullmatch(r"[0-9]{8}", str(only_corp)):
        raise ValueError('DART 법인 고유번호는 8자리 숫자여야 합니다.')

    raw = _db_download_snapshot(cfg)
    _db_valid_snapshot(raw)

    with _db_tempfile.NamedTemporaryFile(suffix='.sqlite3', delete=False) as f:
        f.write(raw)
        tmp = Path(f.name)
    imported = 0
    try:
        with _db_closing(sqlite3.connect(str(tmp))) as remote:
            remote.row_factory = sqlite3.Row
            if not remote.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='company_analysis'"
            ).fetchone():
                return 0
            if only_corp is not None:
                rows = remote.execute(
                    'SELECT corp_code, result_json, analyzed_at FROM company_analysis WHERE corp_code=?',
                    (str(only_corp),),
                ).fetchall()
            else:
                rows = remote.execute(
                    'SELECT corp_code, result_json, analyzed_at FROM company_analysis'
                ).fetchall()
            profiles_exist = bool(remote.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='company_profile'"
            ).fetchone())

            with get_db_connection() as local:
                for row in rows:
                    code = str(row['corp_code'])
                    # 원본 백업 중 비정상 JSON은 무시하고 다른 정상 결과를 계속 병합
                    try:
                        parsed = json.loads(row['result_json'])
                        if not isinstance(parsed, dict) or not parsed.get('company'):
                            continue
                    except (ValueError, TypeError):
                        continue
                    old = local.execute(
                        'SELECT analyzed_at FROM company_analysis WHERE corp_code=?', (code,)
                    ).fetchone()
                    if old and str(old['analyzed_at'] or '') >= str(row['analyzed_at'] or ''):
                        continue
                    local.execute(
                        'INSERT INTO company_analysis(corp_code, result_json, analyzed_at) '
                        'VALUES(?,?,?) ON CONFLICT(corp_code) DO UPDATE SET '
                        'result_json=excluded.result_json, analyzed_at=excluded.analyzed_at',
                        (code, row['result_json'], row['analyzed_at']),
                    )
                    imported += 1
                    if profiles_exist:
                        profile = remote.execute(
                            'SELECT * FROM company_profile WHERE corp_code=?', (code,)
                        ).fetchone()
                        if profile is not None:
                            old_profile = local.execute(
                                'SELECT fetched_at FROM company_profile WHERE corp_code=?', (code,)
                            ).fetchone()
                            if not old_profile or str(old_profile['fetched_at'] or '') < str(profile['fetched_at'] or ''):
                                # 프로필의 구형/신형 DB 간 열 차이를 자동으로 처리
                                local_cols = {x[1] for x in local.execute('PRAGMA table_info(company_profile)')}
                                cols = [key for key in profile.keys() if key in local_cols]
                                placeholders = ','.join('?' for _ in cols)
                                column_list = ','.join('"'+col+'"' for col in cols)
                                values = [profile[key] for key in cols]
                                local.execute(
                                    f'INSERT OR REPLACE INTO company_profile({column_list}) VALUES({placeholders})',
                                    values,
                                )
                # 퇴직연금 영업정보도 기존 원격 SQLite 스냅샷에서 최신 행만 병합.
                meta_exists = bool(remote.execute(
                    "SELECT 1 FROM sqlite_master WHERE type='table' AND name='pension_target_meta'"
                ).fetchone())
                target_imported = 0
                if meta_exists:
                    local.execute('''CREATE TABLE IF NOT EXISTS pension_target_meta (
                        corp_code TEXT PRIMARY KEY,
                        employee_count INTEGER,
                        contact_stage TEXT NOT NULL DEFAULT '미접촉',
                        contact_note TEXT NOT NULL DEFAULT '',
                        updated_at TEXT NOT NULL DEFAULT ''
                    )''')
                    if only_corp is None:
                        meta_rows = remote.execute('SELECT * FROM pension_target_meta').fetchall()
                    else:
                        meta_rows = remote.execute(
                            'SELECT * FROM pension_target_meta WHERE corp_code=?', (str(only_corp),)
                        ).fetchall()
                    for info in meta_rows:
                        code = str(info['corp_code'])
                        if not re.fullmatch(r'\d{8}', code):
                            continue
                        old = local.execute(
                            'SELECT updated_at FROM pension_target_meta WHERE corp_code=?', (code,)
                        ).fetchone()
                        if old and str(old['updated_at'] or '') >= str(info['updated_at'] or ''):
                            continue
                        local.execute('''INSERT INTO pension_target_meta
                             (corp_code, employee_count, contact_stage, contact_note, updated_at)
                             VALUES(?,?,?,?,?) ON CONFLICT(corp_code) DO UPDATE SET
                             employee_count=excluded.employee_count,
                             contact_stage=excluded.contact_stage,
                             contact_note=excluded.contact_note,
                             updated_at=excluded.updated_at''',
                             (code, info['employee_count'], info['contact_stage'],
                              str(info['contact_note'] or '')[:2000], str(info['updated_at'] or '')))
                        target_imported += 1
                st.session_state['_pension_meta_imported'] = target_imported
                # OpenDART 직원현황도 최신 조회 결과만 병합 (기존 수동 입력 테이블은 별도 유지).
                employee_exists = bool(remote.execute(
                    "SELECT 1 FROM sqlite_master WHERE type='table' AND name='pension_employee_dart'"
                ).fetchone())
                if employee_exists:
                    local.execute('''CREATE TABLE IF NOT EXISTS pension_employee_dart (
                       corp_code TEXT PRIMARY KEY, employee_count INTEGER NOT NULL,
                       report_year TEXT NOT NULL, report_code TEXT NOT NULL,
                       report_date TEXT NOT NULL DEFAULT '', report_url TEXT NOT NULL DEFAULT '',
                       checked_at TEXT NOT NULL)''')
                    if only_corp is None:
                        employee_rows=remote.execute('SELECT * FROM pension_employee_dart').fetchall()
                    else:
                        employee_rows=remote.execute('SELECT * FROM pension_employee_dart WHERE corp_code=?',(str(only_corp),)).fetchall()
                    for row in employee_rows:
                        code=str(row['corp_code'])
                        if not re.fullmatch(r'\d{8}',code):
                            continue
                        old=local.execute('SELECT checked_at FROM pension_employee_dart WHERE corp_code=?',(code,)).fetchone()
                        if old and str(old['checked_at'] or '')>=str(row['checked_at'] or ''):
                            continue
                        local.execute('''INSERT INTO pension_employee_dart
                           (corp_code,employee_count,report_year,report_code,report_date,report_url,checked_at)
                           VALUES (?,?,?,?,?,?,?) ON CONFLICT(corp_code) DO UPDATE SET
                           employee_count=excluded.employee_count,report_year=excluded.report_year,
                           report_code=excluded.report_code,report_date=excluded.report_date,
                           report_url=excluded.report_url,checked_at=excluded.checked_at''',
                           (code,row['employee_count'],row['report_year'],row['report_code'],
                            row['report_date'],row['report_url'],row['checked_at']))
                check_exists=bool(remote.execute(
                    "SELECT 1 FROM sqlite_master WHERE type='table' AND name='pension_employee_check'"
                ).fetchone())
                if check_exists:
                    local.execute("""CREATE TABLE IF NOT EXISTS pension_employee_check (
                        corp_code TEXT PRIMARY KEY, status TEXT NOT NULL,
                        detail TEXT NOT NULL DEFAULT '', checked_at TEXT NOT NULL)""")
                    if only_corp is None:
                        check_rows=remote.execute('SELECT * FROM pension_employee_check').fetchall()
                    else:
                        check_rows=remote.execute('SELECT * FROM pension_employee_check WHERE corp_code=?',(str(only_corp),)).fetchall()
                    for row in check_rows:
                        code=str(row['corp_code'])
                        if not re.fullmatch(r'\d{8}',code):
                            continue
                        prior=local.execute('SELECT checked_at FROM pension_employee_check WHERE corp_code=?',(code,)).fetchone()
                        if prior and str(prior['checked_at'] or '')>=str(row['checked_at'] or ''):
                            continue
                        local.execute("""INSERT INTO pension_employee_check
                          (corp_code,status,detail,checked_at) VALUES (?,?,?,?)
                          ON CONFLICT(corp_code) DO UPDATE SET status=excluded.status,
                          detail=excluded.detail,checked_at=excluded.checked_at""",
                          (code,row['status'],row['detail'],row['checked_at']))
        return imported
    finally:
        tmp.unlink(missing_ok=True)


def _db_cloud_restore_on_boot():
    # Once per Python process, before SQLite access. All browser sessions share it.
    cfg = _db_cloud_config()
    if cfg is None:
        return 'not_configured'
    try:
        raw = _db_download_snapshot(cfg, allow_missing=True)
        if raw is None:
            # Storage is empty: seed it with a preexisting, valid DB if present.
            if DB_PATH.exists():
                try:
                    _db_valid_snapshot(_db_sqlite_snapshot())
                    if _db_cloud_push():
                        return 'seeded'
                except Exception:
                    pass
            return 'remote_empty'
        _db_install_snapshot(raw)
        return 'restored'
    except Exception as exc:
        st.session_state['_db_cloud_error'] = '초기 DB 동기화 실패: ' + str(exc)
        return 'offline_fallback'


@st.cache_resource(show_spinner=False)
def _db_boot_once(bucket_object, config_fingerprint):
    # No key in cache key: fingerprint only, not the secret. Cache is per server.
    return _db_cloud_restore_on_boot()


def _db_has_usable_local_master() -> bool:
    """기존 로컬 DB의 회사목록이 있으면 원격 복원 때문에 화면을 멈추지 않는다.

    app.py만 새 폴더에서 실행하면 옛 폴더의 DB는 자동으로 따라오지 않는다.
    그 경우에는 Supabase 복원을 시도하고, 실패해도 복구 UI를 표시한다.
    """
    if not DB_PATH.is_file():
        return False
    try:
        with _db_closing(sqlite3.connect(str(DB_PATH), timeout=2)) as conn:
            exists = conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='corp_master'"
            ).fetchone()
            if not exists:
                return False
            return conn.execute("SELECT EXISTS(SELECT 1 FROM corp_master LIMIT 1)").fetchone()[0] == 1
    except (sqlite3.Error, OSError):
        return False


def _db_startup_restore():
    try:
        cfg = _db_cloud_config()
        fingerprint = _db_hashlib.sha256(
            (cfg['read_url'] + cfg['headers']['apikey']).encode('utf-8')
        ).hexdigest() if cfg else 'disabled'
        return _db_boot_once(_DB_CLOUD_OBJECT, fingerprint)
    except Exception as exc:
        st.session_state['_db_cloud_error'] = str(exc)
        return 'invalid_config'


def _db_cloud_indicator():
    try:
        cfg = _db_cloud_config()
    except Exception as exc:
        st.error(f'클라우드 DB 설정 오류: {exc}')
        return
    if cfg is None:
        st.warning('영구 DB 동기화 미설정: SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY가 필요합니다. 현재 SQLite는 재배포 시 사라질 수 있습니다.')
    else:
        if st.session_state.get('_db_cloud_error'):
            st.warning('☁️ 최근 동기화 오류: ' + st.session_state['_db_cloud_error'])
        else:
            st.caption('☁️ Supabase 설정 확인 · 마지막 백업 성공 여부는 저장 버튼 결과를 확인하세요')
        if st.session_state.get('_db_cloud_saved_at'):
            st.caption('마지막 클라우드 백업: ' + st.session_state['_db_cloud_saved_at'])
        sync_left, sync_right = st.columns(2)
        with sync_left:
            if st.button('☁️ 현재 DB 영구 저장 (압축)', key='corp_cloud_save_btn'):
                _db_cloud_push(show_feedback=True)
        with sync_right:
            if st.button('📥 PC 최신 분석·영업정보 가져오기', key='corp_cloud_pull_analysis_btn'):
                try:
                    with st.spinner('Supabase DB에서 PC가 분석한 결과를 확인하고 병합하는 중입니다...'):
                        count = _db_cloud_merge_analysis_from_remote()
                    meta_count = st.session_state.pop('_pension_meta_imported', 0)
                    if count or meta_count:
                        st.success(f'신규/갱신 분석 {count:,}건 · 퇴직연금 영업정보 {meta_count:,}건 가져왔습니다.')
                        st.rerun()
                    else:
                        st.info('가져올 신규 분석·퇴직연금 영업정보가 없습니다. PC에서 Supabase 저장 여부를 확인해 주세요.')
                except Exception as exc:
                    st.error(f'최신 분석 결과 가져오기 실패: {_dart_safe_error(exc) if isinstance(exc, requests.exceptions.RequestException) else exc}')


def sync_corp_master(corp_df: pd.DataFrame):
    incoming_count = int(len(corp_df))
    incoming_max_date = str(corp_df["modify_date"].max()) if incoming_count and "modify_date" in corp_df.columns else ""
    with get_db_connection() as conn:
        current_count = int(conn.execute("SELECT COUNT(*) FROM corp_master").fetchone()[0])
        current_max_date = str(conn.execute("SELECT COALESCE(MAX(modify_date),'') FROM corp_master").fetchone()[0] or "")
        if current_count == incoming_count and current_max_date == incoming_max_date:
            return

    rows = [
        (str(r.corp_code), str(r.corp_name), str(r.stock_code), str(r.modify_date), str(r.norm_name))
        for r in corp_df[["corp_code", "corp_name", "stock_code", "modify_date", "norm_name"]].itertuples(index=False)
    ]
    with get_db_connection() as conn:
        conn.executemany(
            """
            INSERT INTO corp_master(corp_code, corp_name, stock_code, modify_date, norm_name)
            VALUES(?,?,?,?,?)
            ON CONFLICT(corp_code) DO UPDATE SET
                corp_name=excluded.corp_name,
                stock_code=excluded.stock_code,
                modify_date=excluded.modify_date,
                norm_name=excluded.norm_name
            """,
            rows,
        )


def address_db_stats():
    with get_db_connection() as conn:
        total = conn.execute("SELECT COUNT(*) FROM corp_master").fetchone()[0]
        done = conn.execute(
            "SELECT COUNT(*) FROM company_profile WHERE COALESCE(adres,'') <> ''"
        ).fetchone()[0]
        failed = conn.execute(
            "SELECT COUNT(*) FROM company_profile WHERE COALESCE(fetch_error,'') <> ''"
        ).fetchone()[0]
    return int(total), int(done), int(failed)


def get_profile_rows_for_search(terms, limit=1000):
    terms = [str(x).strip() for x in terms if str(x).strip()]
    sql = """
        SELECT p.corp_code, p.corp_name, p.stock_code, p.ceo_nm, p.corp_cls,
               p.jurir_no, p.bizr_no, p.adres, p.hm_url, p.phn_no, p.induty_code,
               p.modify_date, p.fetched_at
        FROM company_profile p
        WHERE COALESCE(p.adres,'') <> ''
    """
    params = []
    for term in terms:
        sql += " AND p.adres LIKE ?"
        params.append(f"%{term}%")
    sql += " ORDER BY p.corp_name LIMIT ?"
    params.append(int(limit))
    with get_db_connection() as conn:
        rows = conn.execute(sql, params).fetchall()
    return pd.DataFrame([dict(r) for r in rows])


def get_missing_profiles(batch_size=100, listed_first=True, retry_errors=False):
    where = "p.corp_code IS NULL"
    if retry_errors:
        where = "p.corp_code IS NULL OR COALESCE(p.fetch_error,'') <> ''"
    order = "CASE WHEN COALESCE(m.stock_code,'') <> '' THEN 0 ELSE 1 END, m.corp_name" if listed_first else "m.corp_name"
    sql = f"""
        SELECT m.corp_code, m.corp_name, m.stock_code, m.modify_date
        FROM corp_master m
        LEFT JOIN company_profile p ON p.corp_code = m.corp_code
        WHERE {where}
        ORDER BY {order}
        LIMIT ?
    """
    with get_db_connection() as conn:
        rows = conn.execute(sql, (int(batch_size),)).fetchall()
    return [dict(r) for r in rows]


def save_company_profile(master_row: dict, company: dict | None = None, error: str = "", *, persist: bool = True):
    company = company or {}
    values = (
        master_row.get("corp_code", ""),
        company.get("corp_name") or master_row.get("corp_name", ""),
        company.get("stock_code") or master_row.get("stock_code", ""),
        master_row.get("modify_date", ""),
        company.get("ceo_nm", ""),
        company.get("corp_cls", ""),
        company.get("jurir_no", ""),
        company.get("bizr_no", ""),
        company.get("adres", ""),
        company.get("hm_url", ""),
        company.get("ir_url", ""),
        company.get("phn_no", ""),
        company.get("fax_no", ""),
        company.get("induty_code", ""),
        datetime.now().isoformat(timespec="seconds"),
        str(error)[:500],
    )
    with get_db_connection() as conn:
        conn.execute(
            """
            INSERT INTO company_profile(
                corp_code, corp_name, stock_code, modify_date, ceo_nm, corp_cls,
                jurir_no, bizr_no, adres, hm_url, ir_url, phn_no, fax_no,
                induty_code, fetched_at, fetch_error
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(corp_code) DO UPDATE SET
                corp_name=excluded.corp_name,
                stock_code=excluded.stock_code,
                modify_date=excluded.modify_date,
                ceo_nm=excluded.ceo_nm,
                corp_cls=excluded.corp_cls,
                jurir_no=excluded.jurir_no,
                bizr_no=excluded.bizr_no,
                adres=excluded.adres,
                hm_url=excluded.hm_url,
                ir_url=excluded.ir_url,
                phn_no=excluded.phn_no,
                fax_no=excluded.fax_no,
                induty_code=excluded.induty_code,
                fetched_at=excluded.fetched_at,
                fetch_error=excluded.fetch_error
            """,
            values,
        )
    if persist:
        _db_cloud_push()


def fetch_company_info_uncached(api_key: str, corp_code: str) -> dict:
    return request_json(
        f"{DART_BASE}/company.json",
        {"crtfc_key": api_key, "corp_code": corp_code},
        timeout=30,
    )


def build_address_cache(api_key: str, batch_size=100, listed_first=True, retry_errors=False):
    targets = get_missing_profiles(batch_size, listed_first, retry_errors)
    if not targets:
        return 0, 0

    progress = st.progress(0)
    status_box = st.empty()
    ok = 0
    fail = 0
    for idx, row in enumerate(targets, start=1):
        status_box.write(f"주소 DB 구축 {idx}/{len(targets)} · {row['corp_name']}")
        try:
            info = fetch_company_info_uncached(api_key, row["corp_code"])
            save_company_profile(row, info, "", persist=False)
            ok += 1
        except Exception as exc:
            msg = str(exc)
            save_company_profile(row, {}, msg, persist=False)
            fail += 1
            if "요청 제한" in msg or "020" in msg:
                status_box.warning("OpenDART 요청 한도에 도달해 이번 주소 DB 구축을 중단했습니다. 다음 이용 가능 시점에 이어서 실행하면 됩니다.")
                progress.progress(idx / len(targets))
                if ok or fail:
                    _db_cloud_push()
                return ok, fail
        progress.progress(idx / len(targets))

    status_box.empty()
    if ok or fail:
        _db_cloud_push()
    return ok, fail


def export_address_cache_excel():
    with get_db_connection() as conn:
        df = pd.read_sql_query(
            """
            SELECT corp_code AS DART고유번호, corp_name AS 법인명, stock_code AS 종목코드,
                   ceo_nm AS 대표자, corp_cls AS 법인구분, jurir_no AS 법인등록번호,
                   bizr_no AS 사업자등록번호, adres AS 주소, induty_code AS 업종코드,
                   phn_no AS 전화번호, hm_url AS 홈페이지, modify_date AS DART수정일,
                   fetched_at AS 수집일시, fetch_error AS 오류
            FROM company_profile
            ORDER BY corp_name
            """,
            conn,
        )
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="주소DB")
    out.seek(0)
    return out.getvalue()


def import_address_cache_file(uploaded_file):
    if uploaded_file is None:
        return 0
    raw = uploaded_file.getvalue()
    name = (uploaded_file.name or "").lower()
    if name.endswith((".xlsx", ".xls")):
        df = pd.read_excel(io.BytesIO(raw), dtype=str).fillna("")
    else:
        df = None
        for enc in ("utf-8-sig", "cp949", "euc-kr", "utf-8"):
            try:
                df = pd.read_csv(io.BytesIO(raw), dtype=str, encoding=enc).fillna("")
                break
            except Exception:
                pass
        if df is None:
            raise RuntimeError("CSV 인코딩을 판별하지 못했습니다.")

    rename = {
        "DART고유번호": "corp_code", "법인명": "corp_name", "종목코드": "stock_code",
        "대표자": "ceo_nm", "법인구분": "corp_cls", "법인등록번호": "jurir_no",
        "사업자등록번호": "bizr_no", "주소": "adres", "업종코드": "induty_code",
        "전화번호": "phn_no", "홈페이지": "hm_url", "DART수정일": "modify_date",
        "수집일시": "fetched_at", "오류": "fetch_error",
    }
    df = df.rename(columns=rename)
    if "corp_code" not in df.columns:
        raise RuntimeError("주소 DB 파일에 'DART고유번호' 열이 없습니다.")

    count = 0
    with get_db_connection() as conn:
        for _, r in df.iterrows():
            corp_code = str(r.get("corp_code", "")).strip()
            if not corp_code:
                continue
            master = conn.execute(
                "SELECT corp_code, corp_name, stock_code, modify_date FROM corp_master WHERE corp_code=?",
                (corp_code,),
            ).fetchone()
            master_row = dict(master) if master else {
                "corp_code": corp_code,
                "corp_name": str(r.get("corp_name", "")),
                "stock_code": str(r.get("stock_code", "")),
                "modify_date": str(r.get("modify_date", "")),
            }
            company = {
                "corp_name": str(r.get("corp_name", "")), "stock_code": str(r.get("stock_code", "")),
                "ceo_nm": str(r.get("ceo_nm", "")), "corp_cls": str(r.get("corp_cls", "")),
                "jurir_no": str(r.get("jurir_no", "")), "bizr_no": str(r.get("bizr_no", "")),
                "adres": str(r.get("adres", "")), "hm_url": str(r.get("hm_url", "")),
                "phn_no": str(r.get("phn_no", "")), "induty_code": str(r.get("induty_code", "")),
            }
            save_company_profile(master_row, company, str(r.get("fetch_error", "")), persist=False)
            count += 1
    return count


def _normalize_colname(x):
    return re.sub(r"[\s_\-·./()]+", "", str(x)).lower()


def _find_external_column(df, candidates):
    norm_map = {_normalize_colname(c): c for c in df.columns}
    for cand in candidates:
        key = _normalize_colname(cand)
        if key in norm_map:
            return norm_map[key]
    for cand in candidates:
        key = _normalize_colname(cand)
        for norm, orig in norm_map.items():
            if key and (key in norm or norm in key):
                return orig
    return None


def _external_header_score(values):
    """외부 기업목록에서 실제 헤더 행을 찾기 위한 점수.

    충북 제조업체총람처럼 상단에 제목/공백 행이 있고 실제 열 이름이 3~10행 아래에
    있는 파일을 자동 인식한다. '업 체 명'처럼 공백이 들어간 열 이름도 정규화해서 처리한다.
    """
    cells = [_normalize_colname(v) for v in values if str(v).strip() and str(v).lower() != "nan"]
    if not cells:
        return -1

    name_keys = {_normalize_colname(x) for x in EXTERNAL_COL_CANDIDATES["법인명"]}
    addr_keys = {_normalize_colname(x) for x in EXTERNAL_COL_CANDIDATES["주소"]}
    region_keys = {_normalize_colname(x) for x in EXTERNAL_COL_CANDIDATES["지역"]}
    product_keys = {_normalize_colname(x) for x in EXTERNAL_COL_CANDIDATES["업종"]}
    phone_keys = {_normalize_colname(x) for x in EXTERNAL_COL_CANDIDATES["전화번호"]}

    # 시/군명처럼 자료마다 조금씩 달라지는 열 이름도 헤더 단서로 사용
    extra_region_keys = {_normalize_colname(x) for x in ["시군명", "시군", "시군구", "지역명"]}
    region_keys |= extra_region_keys

    has_name = any(c in name_keys for c in cells)
    has_addr = any(c in addr_keys for c in cells)
    has_region = any(c in region_keys for c in cells)

    # 회사명과 주소(또는 지역)가 같이 있는 행만 실제 헤더 후보로 인정
    if not has_name or not (has_addr or has_region):
        return -1

    score = 10
    score += 5 if has_addr else 0
    score += 3 if has_region else 0
    score += 2 if any(c in product_keys for c in cells) else 0
    score += 2 if any(c in phone_keys for c in cells) else 0
    score += min(len(cells), 10) * 0.1
    return score


def _read_external_table(uploaded_file):
    """CSV/Excel 기업목록을 읽고 실제 헤더 행을 자동 탐지한다.

    반환값: (DataFrame, meta)
    meta에는 파일/시트/헤더행 정보가 들어간다.
    """
    raw = uploaded_file.getvalue()
    name = (uploaded_file.name or "").lower()

    if name.endswith((".xlsx", ".xls")):
        excel = pd.ExcelFile(io.BytesIO(raw))
        best = None
        for sheet in excel.sheet_names:
            try:
                preview = pd.read_excel(
                    excel,
                    sheet_name=sheet,
                    header=None,
                    dtype=str,
                    nrows=60,
                ).fillna("")
            except Exception:
                continue

            for idx, row in preview.iterrows():
                score = _external_header_score(row.tolist())
                if score >= 0 and (best is None or score > best[0]):
                    best = (score, sheet, int(idx))

        if best is not None:
            _, sheet, header_row = best
            df = pd.read_excel(
                excel,
                sheet_name=sheet,
                header=header_row,
                dtype=str,
            ).fillna("")
            meta = {
                "파일": uploaded_file.name,
                "시트": sheet,
                "헤더행": header_row + 1,
                "헤더자동감지": True,
            }
            return df, meta

        # 일반적인 1행 헤더 파일은 기존 방식으로도 처리
        df = pd.read_excel(excel, sheet_name=excel.sheet_names[0], dtype=str).fillna("")
        return df, {
            "파일": uploaded_file.name,
            "시트": excel.sheet_names[0],
            "헤더행": 1,
            "헤더자동감지": False,
        }

    # CSV도 상단에 제목 행이 있을 수 있으므로 먼저 header=None으로 읽어 헤더를 탐색
    last_error = None
    for enc in ("utf-8-sig", "cp949", "euc-kr", "utf-8"):
        try:
            preview = pd.read_csv(
                io.BytesIO(raw),
                dtype=str,
                encoding=enc,
                header=None,
                nrows=60,
                engine="python",
            ).fillna("")
            best_row = None
            best_score = -1
            for idx, row in preview.iterrows():
                score = _external_header_score(row.tolist())
                if score > best_score:
                    best_score = score
                    best_row = int(idx)

            if best_score >= 0 and best_row is not None:
                df = pd.read_csv(
                    io.BytesIO(raw),
                    dtype=str,
                    encoding=enc,
                    header=best_row,
                    engine="python",
                ).fillna("")
                return df, {
                    "파일": uploaded_file.name,
                    "인코딩": enc,
                    "헤더행": best_row + 1,
                    "헤더자동감지": True,
                }

            df = pd.read_csv(
                io.BytesIO(raw),
                dtype=str,
                encoding=enc,
                engine="python",
            ).fillna("")
            return df, {
                "파일": uploaded_file.name,
                "인코딩": enc,
                "헤더행": 1,
                "헤더자동감지": False,
            }
        except Exception as e:
            last_error = e
            continue

    raise RuntimeError(f"CSV를 읽지 못했습니다: {last_error or '인코딩/형식을 확인해 주세요.'}")


def read_external_company_file(uploaded_file):
    df, read_meta = _read_external_table(uploaded_file)

    # 완전히 빈 열/행 정리
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all").fillna("")
    df.columns = [str(c).strip() for c in df.columns]

    mapped = {}
    for canonical, candidates in EXTERNAL_COL_CANDIDATES.items():
        mapped[canonical] = _find_external_column(df, candidates)

    if not mapped["법인명"]:
        cols = ", ".join([str(c) for c in df.columns[:20]])
        raise RuntimeError(
            "회사명/법인명/업체명에 해당하는 열을 찾지 못했습니다. "
            f"현재 인식된 열: {cols}"
        )
    if not mapped["주소"] and not mapped["지역"]:
        cols = ", ".join([str(c) for c in df.columns[:20]])
        raise RuntimeError(
            "주소/본사주소/공장주소/소재지에 해당하는 열을 찾지 못했습니다. "
            f"현재 인식된 열: {cols}"
        )

    out = pd.DataFrame()
    for canonical in EXTERNAL_COL_CANDIDATES:
        col = mapped.get(canonical)
        out[canonical] = df[col].astype(str).str.strip() if col else ""

    out["검색주소"] = (out["지역"].fillna("") + " " + out["주소"].fillna("")).str.strip()
    out["norm_name"] = out["법인명"].map(normalize_company_name)
    out = out[out["법인명"].astype(str).str.strip() != ""].reset_index(drop=True)

    # UI에서 사용자가 어떤 행/열을 인식했는지 확인할 수 있도록 메타 포함
    mapped_with_meta = dict(mapped)
    mapped_with_meta["__파일인식정보__"] = read_meta
    return out, mapped_with_meta

def merge_external_company_rows(existing: pd.DataFrame, incoming: pd.DataFrame) -> pd.DataFrame:
    """기존 외부 기업목록을 유지하며 신규 회사/사업장만 추가한다.

    동일 법인명이라도 사업장 소재지가 다르면 별도 행으로 보존한다.
    기존에 저장된 행은 수정하거나 삭제하지 않는다.
    """
    if incoming is None or incoming.empty:
        return existing.copy() if isinstance(existing, pd.DataFrame) else pd.DataFrame()

    current = existing.copy() if isinstance(existing, pd.DataFrame) else pd.DataFrame()
    cols = list(dict.fromkeys([*current.columns, *incoming.columns]))
    current = current.reindex(columns=cols, fill_value="").fillna("")
    additions = incoming.reindex(columns=cols, fill_value="").fillna("")

    def make_keys(df: pd.DataFrame) -> pd.Series:
        names = df["norm_name"].astype(str).str.strip().str.lower()
        locations = (
            df["검색주소"].astype(str)
            .str.replace(r"\s+", "", regex=True).str.lower()
        )
        return names + "|" + locations

    incoming_keys = make_keys(additions)
    known_keys = set(make_keys(current)) if not current.empty else set()
    unique_new = additions.loc[
        ~incoming_keys.isin(known_keys) & ~incoming_keys.duplicated()
    ]
    if unique_new.empty:
        return current
    return pd.concat([current, unique_new], ignore_index=True)


def filter_external_by_region(df: pd.DataFrame, terms, name_query="", limit=3000):
    if df is None or df.empty:
        return pd.DataFrame()
    x = df.copy()
    series = x["검색주소"].fillna("").astype(str)
    for term in [str(t).strip() for t in terms if str(t).strip()]:
        x = x[series.loc[x.index].str.contains(re.escape(term), case=False, na=False)]
    q = normalize_company_name(name_query)
    if q:
        x = x[x["norm_name"].str.contains(re.escape(q), na=False)]
    return x.head(int(limit)).copy()


def dart_match_for_name(corp_df: pd.DataFrame, company_name: str):
    q = normalize_company_name(company_name)
    if not q:
        return []
    exact = corp_df[corp_df["norm_name"] == q]
    if not exact.empty:
        return exact.head(20).to_dict("records")
    contains = corp_df[corp_df["norm_name"].str.contains(re.escape(q), na=False)]
    return contains.head(20).to_dict("records")


def region_terms_from_inputs(province, city, district, town):
    terms = []
    if province and province != "전체":
        terms.append(province)
    for x in (city, district, town):
        x = str(x or "").strip()
        if x and x != "전체":
            terms.append(x)
    return terms


# =========================================================
# DART API
# =========================================================
@st.cache_data(ttl=86400, show_spinner=False)
def load_corp_codes(api_key: str) -> pd.DataFrame:
    url = f"{DART_BASE}/corpCode.xml"
    r = dart_get(url, params={"crtfc_key": api_key}, timeout=60)
    _check_dart_http(r)

    try:
        zf = zipfile.ZipFile(io.BytesIO(r.content))
    except zipfile.BadZipFile:
        msg = decode_bytes(r.content)
        raise RuntimeError(f"고유번호 파일을 열 수 없습니다. API 키를 확인해 주세요.\n{msg[:300]}")

    names = zf.namelist()
    xml_name = next((n for n in names if n.lower().endswith(".xml")), None)
    if not xml_name:
        raise RuntimeError("DART 고유번호 XML 파일을 찾지 못했습니다.")

    xml_text = decode_bytes(zf.read(xml_name))
    root = ET.fromstring(xml_text)

    rows = []
    for item in root.findall(".//list"):
        def _t(tag):
            node = item.find(tag)
            return (node.text or "").strip() if node is not None else ""

        rows.append({
            "corp_code": _t("corp_code"),
            "corp_name": _t("corp_name"),
            "stock_code": _t("stock_code"),
            "modify_date": _t("modify_date"),
        })

    df = pd.DataFrame(rows)
    if df.empty:
        raise RuntimeError("DART 법인 고유번호 목록이 비어 있습니다.")
    df["norm_name"] = df["corp_name"].map(normalize_company_name)
    return df


@st.cache_data(ttl=3600, show_spinner=False)
def get_company_info(api_key: str, corp_code: str) -> dict:
    return request_json(
        f"{DART_BASE}/company.json",
        {"crtfc_key": api_key, "corp_code": corp_code},
        timeout=30,
    )


@st.cache_data(ttl=1800, show_spinner=False)
def get_disclosures(api_key: str, corp_code: str, years_back=5) -> pd.DataFrame:
    end = date.today()
    start = end - timedelta(days=365 * years_back + 30)

    params = {
        "crtfc_key": api_key,
        "corp_code": corp_code,
        "bgn_de": start.strftime("%Y%m%d"),
        "end_de": end.strftime("%Y%m%d"),
        "page_no": 1,
        "page_count": 100,
    }
    data = request_json(f"{DART_BASE}/list.json", params, timeout=30)
    items = data.get("list", [])
    df = pd.DataFrame(items)
    return df


def pick_latest_audit_report(disclosures: pd.DataFrame):
    if disclosures is None or disclosures.empty or "report_nm" not in disclosures.columns:
        return None

    df = disclosures.copy()
    # 국문/영문 OpenDART 모두 지원
    audit_pattern = r"감사보고서|audit\s*report|auditor['’]?s?\s*report|independent\s*auditor"
    mask = df["report_nm"].astype(str).str.contains(
        audit_pattern, case=False, na=False, regex=True
    )
    audit = df[mask].copy()

    if audit.empty:
        return None

    # 최근 접수일 / 접수번호 기준
    sort_cols = [c for c in ["rcept_dt", "rcept_no"] if c in audit.columns]
    if sort_cols:
        audit = audit.sort_values(sort_cols, ascending=False)
    return audit.iloc[0].to_dict()


@st.cache_data(ttl=3600, show_spinner=False)
def download_disclosure_text(api_key: str, rcept_no: str) -> str:
    url = f"{DART_BASE}/document.xml"
    r = dart_get(
        url,
        params={"crtfc_key": api_key, "rcept_no": rcept_no},
        timeout=90,
    )
    _check_dart_http(r)

    try:
        zf = zipfile.ZipFile(io.BytesIO(r.content))
    except zipfile.BadZipFile:
        msg = decode_bytes(r.content)
        raise RuntimeError(f"공시 원문 ZIP을 열 수 없습니다.\n{msg[:400]}")

    candidates = []
    for name in zf.namelist():
        lower = name.lower()
        if lower.endswith((".xml", ".html", ".htm", ".txt")):
            info = zf.getinfo(name)
            # 비정상적으로 큰 단일 파일은 메모리 과다 사용 방지
            if info.file_size <= 20 * 1024 * 1024:
                candidates.append((name, info.file_size))

    if not candidates:
        raise RuntimeError("분석 가능한 감사보고서 원문 파일을 찾지 못했습니다.")

    # 텍스트 길이가 충분한 문서들을 우선적으로 합침
    parts = []
    total_chars = 0
    for name, _ in sorted(candidates, key=lambda x: x[1], reverse=True):
        raw = zf.read(name)
        txt = clean_text(decode_bytes(raw))
        if len(txt) < 50:
            continue
        parts.append(txt)
        total_chars += len(txt)
        if total_chars > 2_500_000:
            break

    return "\n\n".join(parts)


# =========================================================
# 퇴직연금 / 퇴직급여 공시 분석 (v2.9)
# =========================================================
PENSION_SIGNALS = {
    "DB형": [r"확정급여(?:형|제도|채무|부채)", r"사외적립자산", r"순확정급여", r"defined\s+benefit", r"plan\s+assets"],
    "DC형": [r"확정기여(?:형|제도)", r"defined\s+contribution"],
    "퇴직연금": [r"퇴직연금", r"퇴직급여", r"퇴직급여충당", r"종업원급여", r"employee\s+benefits?", r"retirement\s+benefits?"],
}
PENSION_AMOUNT_PATTERNS = {
    "사외적립자산": [r"사외적립자산", r"적립자산의\s*공정가치", r"fair\s+value\s+of\s+plan\s+assets", r"plan\s+assets"],
    "확정급여채무": [r"확정급여채무", r"확정급여채무의\s*현재가치", r"present\s+value\s+of\s+defined\s+benefit"],
    "순확정급여부채": [r"순확정급여부채", r"net\s+defined\s+benefit\s+liabilit"],
    "DC 부담금/비용": [r"확정기여제도.*(?:비용|납입|부담금)", r"contribution\s+expense"],
}
PENSION_EXTRA_TERMS = [r"퇴직연금사업자", r"퇴직연금운용기관", r"퇴직연금운용자산", r"연금자산", r"연금신탁"]



# 제도 '운영'의 직접 근거와 회계정책상 일반 설명을 구별한다.
# 추정 분류를 운영 확정으로 오인하지 않도록 보수적으로 설정한다.
PENSION_NEGATIVE_RE = re.compile(
    r"(?:퇴직연금(?:제도)?|확정급여제도|확정기여제도).{0,28}(?:운영하지\s*않|가입하지\s*않|도입하지\s*않|미운영|미가입|미도입|없(?:습니|다))|"
    r"(?:퇴직연금(?:제도)?|확정급여제도|확정기여제도)\s*(?:미운영|미가입|미도입)", re.I)
PENSION_DB_DIRECT_RE = re.compile(
    r"(?:당사|연결회사|회사는|회사의|회사는\s*현재|임직원|종업원).{0,90}"
    r"(?:확정급여(?:형|제도)|DB형).{0,80}(?:운영|가입|도입|적용|유지|설정)|"
    r"(?:확정급여(?:형|제도)|DB형).{0,90}(?:운영하고|운용하고|가입하고|도입하였|적용하고|유지하고)|"
    r"(?:당기말|보고기간말|기말).{0,75}(?:사외적립자산|확정급여채무).{0,75}\d[\d,]*|"
    r"(?:사외적립자산|확정급여채무).{0,50}\d[\d,]*(?:\s*(?:천원|백만원|억원|원))", re.I)
PENSION_DC_DIRECT_RE = re.compile(
    r"(?:당사|연결회사|회사는|임직원|종업원).{0,90}"
    r"(?:확정기여(?:형|제도)|DC형).{0,80}(?:운영|가입|도입|적용|유지)|"
    r"(?:확정기여(?:형|제도)|DC형).{0,90}(?:운영하고|운용하고|가입하고|도입하였|적용하고|유지하고)|"
    r"(?:확정기여(?:형|제도)|DC형).{0,90}(?:납입|부담금|기여금|비용).{0,50}\d[\d,]*|"
    r"(?:확정기여(?:형|제도)|DC형).{0,50}\d[\d,]*(?:\s*(?:천원|백만원|억원|원))", re.I)
PENSION_POLICY_RE = re.compile(r"회계정책|회계처리|측정|인식|평가|추정|기업회계기준|공정가치|회계기준|부채로\s*인식|비용으로\s*인식", re.I)


def classify_pension_operations(raw_text: str) -> dict:
    """운영 증거를 우선, 불충분하면 미확인. 회사의 실제 연금계약을 보증하지 않음."""
    blocks = [re.sub(r"\s+", " ", line).strip() for line in re.split(r"\n+", raw_text or "") if line.strip()]
    direct = {"DB": [], "DC": [], "미운용": []}
    clues = {"DB": [], "DC": []}
    for line in blocks:
        snippet = line[:650]
        # 원문에 숫자나 '운영'이 등장하더라도 회계정책을 설명하는 가정문이면 직접 증거 제외
        policy_or_conditional = bool(PENSION_POLICY_RE.search(line) and re.search(r"경우|하여|으로\s*표시|있습니다|가정|회계정책",line,re.I))
        conditional = bool(re.search(r"(?:할\s*경우|하는\s*경우|경우에는|경우\s*해당|가정하여)",line,re.I))
        neg=bool(PENSION_NEGATIVE_RE.search(line))
        if neg and len(direct["미운용"])<4:
            direct["미운용"].append(snippet)
        for key, keyword, pattern in [("DB", r"확정급여|사외적립자산|DB형|defined\s+benefit", PENSION_DB_DIRECT_RE),
                                      ("DC", r"확정기여|DC형|defined\s+contribution", PENSION_DC_DIRECT_RE)]:
            if re.search(keyword,line,re.I) and len(clues[key]) < 5:
                clues[key].append(snippet)
            if not neg and not policy_or_conditional and not conditional and pattern.search(line) and len(direct[key])<5:
                direct[key].append(snippet)
    db,dc = bool(direct["DB"]),bool(direct["DC"])
    non=bool(direct["미운용"])
    if non and (db or dc):
        label="확인 불가"; confidence="상충 근거"
    elif db and dc:
        label="DB+DC 혼합"; confidence="직접 근거 확인"
    elif db:
        label="DB형 운영"; confidence="직접 근거 확인"
    elif dc:
        label="DC형 운영"; confidence="직접 근거 확인"
    elif non:
        label="미운용 확인"; confidence="명시적 부정 공시"
    else:
        label="확인 불가"; confidence="운영 직접 근거 부족"
    return {"운영분류":label,"판별신뢰":confidence,"운영근거":direct,"관련단서":clues}


def analyze_pension_text(raw_text: str) -> dict:
    """원문 공시 내 근거 탐색. 표 단위·금액 열 연계가 확실하지 않으면 임의 금액 추출 금지."""
    blocks = [re.sub(r"\s+", " ", x).strip() for x in re.split(r"\n+", raw_text or "") if x.strip()]
    signals = {k: [] for k in PENSION_SIGNALS}
    amounts = {k: [] for k in PENSION_AMOUNT_PATTERNS}
    institution_evidence = []
    retirement_words = re.compile(r"퇴직연금|사외적립|확정급여|확정기여|연금운용|pension|retirement|plan\s+assets|defined\s+(?:benefit|contribution)", re.I)
    for i, line in enumerate(blocks):
        # 제목, 행/값이 다른 줄에 나뉜 경우 앞뒤 문장을 원문 그대로 첨부
        if not retirement_words.search(line):
            continue
        context = " | ".join(blocks[max(0,i-1):min(len(blocks),i+3)])[:700]
        for category, patterns in PENSION_SIGNALS.items():
            if any(re.search(p, line, re.I) for p in patterns) and context not in signals[category] and len(signals[category])<7:
                signals[category].append(context)
        for field, patterns in PENSION_AMOUNT_PATTERNS.items():
            if any(re.search(p, line, re.I) for p in patterns) and context not in amounts[field] and len(amounts[field])<6:
                amounts[field].append(context)
        # 기관명은 금융기관 이름과 연금 문맥이 같은 좁은 범위에 존재할 때만 '후보'로 표시
        for bank, patterns in BANK_PATTERNS.items():
            if re.search("|".join(f"(?:{v})" for v in patterns), line, re.I):
                if re.search(r"퇴직연금|연금사업자|연금신탁|사외적립|pension|retirement", line, re.I):
                    item={"금융기관":bank,"판정":"연금 관련 문맥에 기관명 병기(역할 확인 필요)","근거":context}
                    if item not in institution_evidence:
                        institution_evidence.append(item)
    db = bool(signals["DB형"])
    dc = bool(signals["DC형"])
    if db and dc:
        status="DB형·DC형 관련 공시 모두 확인"
    elif db:
        status="DB형 관련 공시 확인"
    elif dc:
        status="DC형 관련 공시 확인"
    elif signals["퇴직연금"]:
        status="퇴직급여 관련 공시 확인 · 유형 불명"
    else:
        status="공시 근거 미확인"
    operations = classify_pension_operations(raw_text)
    return {**operations, "판정": status, "DB근거": signals["DB형"], "DC근거": signals["DC형"],
            "일반근거":signals["퇴직연금"], "금액관련근거":amounts,
            "연금기관후보":institution_evidence[:12],
            "안내":"DART 공개문서 기반 단서이며 실제 제도 가입·계좌별 적립금·수탁/운용기관을 확정하지 않습니다. 표의 단위·기준일·열 구성을 원문에서 확인하세요."}


def render_pension_analysis(pension: dict):
    """공시 근거에 기반한 퇴직연금 영업 대시보드. 미검증 금액은 노출하지 않음."""
    st.markdown("### 📊 퇴직연금 영업분석 대시보드")
    st.caption("감사보고서 공시 기반 · 확정 정보와 추가 확인이 필요한 사항을 구분")

    status = pension.get("판정") or "공시 근거 미확인"
    pension_type = pension.get("운영분류", "확인 불가")
    confidence = pension.get("판별신뢰", "운영 직접 근거 부족")
    db = bool(pension.get("DB근거"))
    dc = bool(pension.get("DC근거"))
    institutions = pension.get("연금기관후보") or []
    amt = pension.get("금액관련근거") or {}
    asset_rows = amt.get("사외적립자산") or []
    obligation_rows = amt.get("확정급여채무") or []
    institution_label = f"후보 {len(institutions)}개" if institutions else "미확인"

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.metric("퇴직연금 운영 분류", pension_type)
        st.caption(confidence)
    with c2:
        st.metric("사외적립자산", "금액 미검증")
        st.caption(f"관련 원문 {len(asset_rows)}건" if asset_rows else "금액 관련 원문 없음")
    with c3:
        st.metric("확정급여채무", "금액 미검증")
        st.caption(f"관련 원문 {len(obligation_rows)}건" if obligation_rows else "금액 관련 원문 없음")
    with c4:
        st.metric("퇴직연금 금융기관", institution_label)
        st.caption("기관명·역할 원문 확인 필요")

    with st.container(border=True):
        st.markdown("#### 🎯 영업 검토 요약")
        if pension_type in ("DB형 운영", "DC형 운영", "DB+DC 혼합"):
            st.success(f"**{pension_type}** — 해당 유형의 운영을 시사하는 직접 공시 근거를 발견했습니다. 계약 정보는 별도 확인하세요.")
        elif pension_type == "미운용 확인":
            st.info("퇴직연금 미운용에 관한 명시적인 문구가 있습니다. 기준일과 회사 전체 적용 여부를 원문에서 확인하세요.")
        else:
            st.warning("운영 유형 확인 불가: 관련 문구만으로 운영 또는 미운영을 단정할 수 없습니다.")
        st.caption(f"기존 키워드 분석: {status}")
        op = pension.get("운영근거") or {}
        with st.expander("운영형태 판별 근거 (DB·DC·미운용)", expanded=False):
            for k in ("DB", "DC", "미운용"):
                st.markdown(f"**{k} 직접 근거**")
                for line in op.get(k, []):
                    st.info(line)
                if not op.get(k):
                    st.caption("직접 근거 없음")
        st.markdown("**상담 시 확인할 4가지 사항**")
        st.markdown("① 실제 DB/DC 가입 여부 및 가입 인원  ·  ② 현재 퇴직연금사업자  ·  ③ 적립금·전년도 납입액  ·  ④ 연말 추가납입·사업자 변경 계획")
        st.caption("참고: 사외적립자산/확정급여채무는 법정 최소적립금 충족률을 직접 의미하지 않으며, 공시상 기관명만으로 퇴직연금사업자라고 단정할 수 없습니다.")

    left, right = st.columns([1, 1])
    with left:
        st.markdown("#### 🔍 공시 근거 현황")
        overview = pd.DataFrame([
            {"구분": "DB형", "근거건수": len(pension.get("DB근거") or []), "상태": "관련 문구 있음" if db else "미확인"},
            {"구분": "DC형", "근거건수": len(pension.get("DC근거") or []), "상태": "관련 문구 있음" if dc else "미확인"},
            {"구분": "사외적립자산", "근거건수": len(asset_rows), "상태": "금액 검증 필요" if asset_rows else "미확인"},
            {"구분": "확정급여채무", "근거건수": len(obligation_rows), "상태": "금액 검증 필요" if obligation_rows else "미확인"},
        ])
        st.dataframe(overview, use_container_width=True, hide_index=True)
    with right:
        st.markdown("#### 🏦 금융기관 연결 근거")
        if institutions:
            st.dataframe(pd.DataFrame(institutions), use_container_width=True, hide_index=True)
        else:
            st.info("연금 문맥에서 특정 금융기관을 직접 확인하지 못했습니다.")
        st.caption("기존 거래은행 목록은 연금 운용기관 목록과 별개입니다.")

    st.markdown("#### 📑 감사보고서 상세 근거")
    for label, key in [("DB형 관련 근거", "DB근거"), ("DC형 관련 근거", "DC근거"), ("기타 퇴직급여 근거", "일반근거")]:
        rows = pension.get(key) or []
        if rows:
            with st.expander(f"{label} ({len(rows)}건)", expanded=False):
                for i, row in enumerate(rows, 1):
                    st.markdown(f"**근거 {i}**")
                    st.info(row)
    with st.expander("사외적립자산·급여채무·DC 부담금 원문 (단위/기준일 검증)", expanded=False):
        found = False
        for label, rows in amt.items():
            if rows:
                found = True
                st.markdown(f"**{label}**")
                for row in rows:
                    st.info(row)
        if not found:
            st.info("관련 원문을 발견하지 못했습니다.")
    st.caption(pension.get("안내", ""))


# =========================================================
# 분석
# =========================================================
def bank_hits_in_text(text: str):
    hits = []

    # 문맥 단위: 너무 짧은 줄은 인접 문장으로 합치기 위해 우선 문단 기반
    blocks = [b.strip() for b in re.split(r"\n{1,}", text) if b.strip()]
    blocks = [re.sub(r"\s+", " ", b) for b in blocks]

    for bank, patterns in BANK_PATTERNS.items():
        regex = re.compile("|".join(f"(?:{p})" for p in patterns), re.IGNORECASE)

        contexts = []
        categories = set()
        monies = []

        for i, block in enumerate(blocks):
            if not regex.search(block):
                continue

            # 표/문장 조각이 너무 짧으면 주변 1줄 결합
            pieces = []
            if i > 0:
                pieces.append(blocks[i - 1])
            pieces.append(block)
            if i + 1 < len(blocks):
                pieces.append(blocks[i + 1])

            context = " | ".join(pieces)
            context = context[:800]

            # 관련 거래 키워드가 같이 있는 문맥을 우선
            related = any(word in context for word in RELATED_WORDS)
            if related or len(contexts) < 3:
                if context not in contexts:
                    contexts.append(context)

            for cat, cat_patterns in CATEGORY_PATTERNS.items():
                if any(re.search(p, context, re.IGNORECASE) for p in cat_patterns):
                    categories.add(cat)

            monies.extend(money_tokens(context))

            if len(contexts) >= 8:
                break

        if contexts:
            unique_money = []
            for m in monies:
                if m not in unique_money:
                    unique_money.append(m)

            # "주거래은행" 단정 대신 증빙 강도만 표시
            evidence_score = len(categories) + min(len(contexts), 3)
            if evidence_score >= 5:
                strength = "높음"
            elif evidence_score >= 3:
                strength = "보통"
            else:
                strength = "낮음"

            hits.append({
                "은행": bank,
                "거래유형": ", ".join(sorted(categories)) if categories else "은행명 확인",
                "확인강도": strength,
                "금액표기": ", ".join(unique_money[:4]),
                "근거건수": len(contexts),
                "근거문구": "\n\n".join(contexts[:5]),
            })

    # 은행명이 특정되지 않았지만 금융기관 관련 문구가 있는 경우
    generic_contexts = []
    generic_regex = re.compile("|".join(map(re.escape, RELATED_WORDS)))
    for block in blocks:
        if generic_regex.search(block):
            if len(block) > 15:
                generic_contexts.append(block[:600])
        if len(generic_contexts) >= 10:
            break

    return hits, generic_contexts


def search_companies(corp_df: pd.DataFrame, query: str) -> pd.DataFrame:
    q = normalize_company_name(query)
    if not q:
        return corp_df.iloc[0:0].copy()

    exact = corp_df[corp_df["norm_name"] == q]
    if not exact.empty:
        return exact.head(20).copy()

    contains = corp_df[corp_df["norm_name"].str.contains(re.escape(q), na=False)]
    return contains.head(30).copy()


def company_label(row):
    stock = row.get("stock_code", "")
    stock_txt = f" / 종목코드 {stock}" if stock else ""
    return f'{row["corp_name"]} / DART {row["corp_code"]}{stock_txt}'


def _analyze_company_online(api_key: str, corp_row: dict):
    corp_code = corp_row["corp_code"]
    corp_name = corp_row["corp_name"]

    # 저장된 기업개황을 이용한다. company.json은 부가정보이므로
    # DART 연결이 불안정할 때 감사보고서 분석을 막지 않도록 강제 재호출하지 않는다.
    company = get_saved_company_profile(str(corp_code)) or {}
    try:
        disclosures = get_disclosures(api_key, corp_code)
    except Exception as exc:
        if isinstance(exc, (RuntimeError, requests.exceptions.RequestException)):
            raise RuntimeError(f"DART 공시목록(list.json) 조회 실패: {exc if isinstance(exc, RuntimeError) else _dart_safe_error(exc)}") from None
        raise
    report = pick_latest_audit_report(disclosures)

    base = {
        "법인명": corp_name,
        "DART고유번호": corp_code,
        "대표자": company.get("ceo_nm", ""),
        "사업자등록번호": company.get("bizr_no", ""),
        "법인등록번호": company.get("jurir_no", ""),
        "주소": company.get("adres", ""),
        "전화번호": company.get("phn_no", ""),
        "홈페이지": company.get("hm_url", ""),
    }

    if not report:
        return {
            "company": base,
            "report": None,
            "banks": [],
            "generic": [],
            "pension": analyze_pension_text(""),
            "error": "최근 5년 공시에서 '감사보고서'를 찾지 못했습니다.",
        }

    rcept_no = str(report.get("rcept_no", ""))
    try:
        text = download_disclosure_text(api_key, rcept_no)
    except Exception as exc:
        if isinstance(exc, (RuntimeError, requests.exceptions.RequestException)):
            raise RuntimeError(f"DART 감사보고서 원문(document.xml) 다운로드 실패: {exc if isinstance(exc, RuntimeError) else _dart_safe_error(exc)}") from None
        raise
    banks, generic = bank_hits_in_text(text)
    pension = analyze_pension_text(text)

    report_info = {
        "보고서명": report.get("report_nm", ""),
        "접수일": report.get("rcept_dt", ""),
        "접수번호": rcept_no,
        "DART링크": f"https://dart.fss.or.kr/dsaf001/main.do?rcpNo={rcept_no}",
    }

    return {
        "company": base,
        "report": report_info,
        "banks": banks,
        "generic": generic,
        "pension": pension,
        "error": "",
    }


def analyze_company(api_key: str, corp_row: dict, *, refresh=False):
    """저장 결과 우선. 사용자 재분석은 새로 요청하되 장애 시 양호한 기존 결과 보호."""
    code = str(corp_row.get("corp_code", ""))
    cached, analyzed_at = get_saved_company_analysis(code)
    if cached is None:
        # 앱을 다시 시작하지 않아도 PC가 Supabase에 올린 최신 분석을 바로 활용.
        # Supabase가 느리거나 실패하면 기존 로컬 캐시 / 직접 DART 조회로 진행한다.
        try:
            _db_cloud_merge_analysis_from_remote(only_corp=code)
            cached, analyzed_at = get_saved_company_analysis(code)
        except Exception:
            pass
    usable_cache = bool(cached and cached.get("report") and not cached.get("error"))

    if not dart_enabled:
        if cached is None:
            raise RuntimeError("저장된 분석 결과가 없습니다. DART 접속 가능할 때 해당 법인을 먼저 분석해 주세요.")
        st.info(f"오프라인: {analyzed_at} 저장 결과 표시 (최신 공시 미반영)")
        return cached

    if usable_cache and not refresh:
        st.info(f"💾 {analyzed_at} 저장 분석 결과를 즉시 표시합니다. 최신 공시는 '최신 공시 재분석'을 눌러 주세요.")
        return cached

    try:
        if refresh:
            # 사용자 명시적 재분석은 장애 이후의 90초 재접속 제한을 해제한다.
            _dart_reset_cooldown()
            get_disclosures.clear(api_key, code)
        result = _analyze_company_online(api_key, corp_row)
        # 잠시 공시 결과가 누락되었거나 영어 경로에서 예상 못한 응답을 받아도
        # 정상적으로 분석된 기존 거래은행/퇴직연금 저장본을 덮어쓰지 않는다.
        if usable_cache and (not result.get("report") or result.get("error")):
            st.warning(
                f"최근 공시에서 유효한 감사보고서를 확인하지 못해 "
                f"{analyzed_at}에 저장한 기존 정상 결과를 유지합니다."
            )
            return cached
        save_company_analysis(result)
        return result
    except Exception as exc:
        if cached is not None:
            reason = str(exc) if isinstance(exc, RuntimeError) else _dart_safe_error(exc)
            st.warning(
                f"DART 최신 조회 실패 ({reason}). "
                f"{analyzed_at}에 저장한 결과를 표시합니다. 최신 공시가 아닐 수 있습니다."
            )
            return cached
        raise


def make_excel(results):
    output = io.BytesIO()

    summary_rows = []
    bank_rows = []
    evidence_rows = []
    pension_rows = []
    pension_evidence_rows = []

    for result in results:
        c = result["company"]
        report = result.get("report") or {}

        summary_rows.append({
            **c,
            "감사보고서": report.get("보고서명", ""),
            "접수일": report.get("접수일", ""),
            "DART링크": report.get("DART링크", ""),
            "확인은행수": len(result.get("banks", [])),
            "오류/비고": result.get("error", ""),
        })

        pen=result.get("pension") or {}
        pension_rows.append({"법인명":c.get("법인명", ""), "운영형태":pen.get("운영분류", "확인 불가"), "판별신뢰":pen.get("판별신뢰", ""), "퇴직연금판정":pen.get("판정", "분석자료 없음"),
                             "연금기관명후보":", ".join(x.get("금융기관", "") for x in pen.get("연금기관후보",[])),
                             "DART링크":report.get("DART링크", "")})
        for category, evidences in (pen.get("운영근거") or {}).items():
            for evidence in evidences:
                pension_evidence_rows.append({"법인명":c.get("법인명", ""),"항목":"운영 직접근거 - "+category,"근거원문":evidence})
        for category, key in [("DB형", "DB근거"),("DC형", "DC근거"),("퇴직급여", "일반근거")]:
            for evidence in pen.get(key, []):
                pension_evidence_rows.append({"법인명":c.get("법인명", ""),"항목":category,"근거원문":evidence})
        for category, evidences in (pen.get("금액관련근거") or {}).items():
            for evidence in evidences:
                pension_evidence_rows.append({"법인명":c.get("법인명", ""),"항목":category,"근거원문":evidence})

        for b in result.get("banks", []):
            bank_rows.append({
                "법인명": c["법인명"],
                "은행": b["은행"],
                "거래유형": b["거래유형"],
                "확인강도": b["확인강도"],
                "금액표기": b["금액표기"],
                "근거건수": b["근거건수"],
                "DART링크": report.get("DART링크", ""),
            })
            evidence_rows.append({
                "법인명": c["법인명"],
                "은행": b["은행"],
                "근거문구": b["근거문구"],
            })

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(summary_rows).to_excel(writer, index=False, sheet_name="법인요약")
        pd.DataFrame(bank_rows).to_excel(writer, index=False, sheet_name="거래은행")
        pd.DataFrame(evidence_rows).to_excel(writer, index=False, sheet_name="근거문구")
        pd.DataFrame(pension_rows).to_excel(writer, index=False, sheet_name="퇴직연금요약")
        pd.DataFrame(pension_evidence_rows, columns=["법인명","항목","근거원문"]).to_excel(writer, index=False, sheet_name="연금근거")

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            # 과도하게 넓은 셀 방지
            for col in ws.columns:
                letter = col[0].column_letter
                max_len = 0
                for cell in col[:100]:
                    value = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, min(len(value), 60))
                ws.column_dimensions[letter].width = max(12, min(max_len + 2, 60))

    output.seek(0)
    return output.getvalue()



# =========================================================
# A4 현장 방문용 요약 보고서 (PDF + 편집가능 Excel)
# 분석결과를 표시할 때만 메모리에서 생성. DB/DART/Supabase 변경 없음.
# =========================================================
def _visit_text(value, fallback='-'):
    """DART/외부기업목록의 원문을 보고서에 안전한 한 줄로 정리."""
    if value is None:
        return fallback
    v = re.sub(r'\s+', ' ', str(value)).strip()
    return v if v and v.lower() not in ('nan', 'none') else fallback


def _visit_cut(s, limit=36):
    s = _visit_text(s)
    return s if len(s) <= limit else s[:limit-1].rstrip() + '…'


def _visit_filename(name):
    # Windows/모바일 파일 저장 시 사용할 수 없는 문자를 제거.
    value = re.sub(r'[\\/:*?"<>|]+', '_', _visit_text(name, '법인'))
    return (value.strip(' .') or '법인')[:60]


def _visit_report_context(result, external_meta=None):
    c = result.get('company') or {}
    rep = result.get('report') or {}
    pen = result.get('pension') or {}
    ext = external_meta or {}
    banks = list(result.get('banks') or [])
    # 공시에서 '차입/대출' 문맥이 있는 금융기관을 우선 표시한다.
    banks.sort(key=lambda b: (0 if '차입/대출' in str(b.get('거래유형') or '') else 1,
                              str(b.get('은행') or '')))
    bank_rows=[]
    for b in banks:
        bank_rows.append({
            'bank': _visit_text(b.get('은행')),
            'type': _visit_text(b.get('거래유형'), '은행명 확인'),
            'amount': _visit_text(b.get('금액표기'), '미확인'),
            'strength': _visit_text(b.get('확인강도'), '미확인'),
            'evidence': _visit_text(((_evidence_items(b) or [''])[0]), ''),
        })
    inst = [_visit_text(x.get('금융기관'), '') for x in (pen.get('연금기관후보') or [])]
    inst = list(dict.fromkeys(i for i in inst if i))
    report_date=_visit_text(rep.get('접수일'))
    if report_date != '-' and re.fullmatch(r'\d{8}',report_date):
        report_date=f'{report_date[:4]}.{report_date[4:6]}.{report_date[6:]}'
    profile = None
    # 불필요한 DART 요청 없이 저장된 기업개황만 참조한다.
    try:
        code = c.get('DART고유번호')
        if code:
            profile = get_saved_company_profile(str(code))
    except (ValueError, sqlite3.Error, OSError):
        profile = None
    profile=profile or {}
    return {
        'name': _visit_text(c.get('법인명') or ext.get('법인명')),
        'ceo': _visit_text(c.get('대표자') or ext.get('대표자') or profile.get('ceo_nm')),
        'phone': _visit_text(c.get('전화번호') or ext.get('전화번호') or profile.get('phn_no')),
        'address': _visit_text(c.get('주소') or ext.get('주소') or profile.get('adres')),
        'bizno': _visit_text(c.get('사업자등록번호') or ext.get('사업자등록번호')),
        'jurino': _visit_text(c.get('법인등록번호') or ext.get('법인등록번호')),
        'industry': _visit_text(ext.get('업종') or profile.get('induty_code'), '미확인'),
        'dart_code': _visit_text(c.get('DART고유번호')),
        'report_name': _visit_text(rep.get('보고서명'), '감사보고서 미확인'),
        'report_date': report_date,
        'report_url': _visit_text(rep.get('DART링크'), ''),
        'generated': datetime.now().strftime('%Y.%m.%d'),
        'banks': bank_rows,
        'pension_type': _visit_text(pen.get('운영분류'), '확인 불가'),
        'pension_confidence': _visit_text(pen.get('판별신뢰'), '판별 근거 없음'),
        'pension_status': _visit_text(pen.get('판정'), '공시 근거 미확인'),
        'pension_inst': ', '.join(inst) if inst else '퇴직연금 기관 미확인',
        'db_count': len(pen.get('DB근거') or []),
        'dc_count': len(pen.get('DC근거') or []),
        'asset_count': len((pen.get('금액관련근거') or {}).get('사외적립자산') or []),
        'error': _visit_text(result.get('error'), ''),
    }


def make_visit_report_excel(result, external_meta=None):
    """기존 상세 엑셀과 별개의, 인쇄 설정이 적용된 편집가능한 A4 한 장짜리 보고서."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.utils import get_column_letter
    d = _visit_report_context(result, external_meta)
    wb = Workbook()
    ws = wb.active
    ws.title = 'A4현장미팅보고서'
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    for i,w in enumerate([12,12,13,13,13,13,12,12],1):
        ws.column_dimensions[get_column_letter(i)].width = w
    navy='18354F'; teal='008879'; dark='253849'; mid='526679'; light='EDF4F4'; white='FFFFFF'; grey='F2F5F7'; orange='FFF7E8'; line='D8E3E7'
    edge=Side(style='thin', color=line)

    def cellrow(row,text, start=1,end=8, *, height=20, size=10, color=dark,
                bold=False, bg=None, align='left', indent=1, wrap=False):
        ws.merge_cells(start_row=row,start_column=start,end_row=row,end_column=end)
        for col in range(start,end+1):
            x=ws.cell(row,col)
            if bg: x.fill=PatternFill('solid',fgColor=bg)
        t=ws.cell(row,start)
        t.value=str(text)
        t.font=Font(name='맑은 고딕',size=size,bold=bold,color=color)
        t.alignment=Alignment(horizontal=align,vertical='center',indent=indent if align=='left' else 0,wrap_text=wrap)
        ws.row_dimensions[row].height=max(height,ws.row_dimensions[row].height or 0)

    def section(row,text):
        cellrow(row,text,height=23,size=10,bold=True,color=white,bg=navy)

    def dual(row, left_label,left_value,right_label,right_value):
        cellrow(row,left_label,1,2,height=23,bg=grey,size=9,bold=True,color=mid)
        cellrow(row,_visit_cut(left_value,25),3,4,size=9)
        cellrow(row,right_label,5,6,bg=grey,size=9,bold=True,color=mid)
        cellrow(row,_visit_cut(right_value,25),7,8,size=9)

    cellrow(1,'CORPORATE FINDER | 현장 미팅 리포트',height=24,size=10,bold=True,color=teal)
    cellrow(2,'법인 분석 요약 보고서',height=35,size=19,bold=True,color=white,bg=navy)
    cellrow(3,f"{d['name']}  |  생성일 {d['generated']}  |  감사보고서 접수일 {d['report_date']}",height=24,size=10,bold=True,color=dark,bg=light)
    ws.row_dimensions[4].height=7
    section(5,'01  기업 기본정보')
    dual(6,'대표자',d['ceo'],'전화번호',d['phone'])
    dual(7,'사업자번호',d['bizno'],'법인번호',d['jurino'])
    cellrow(8,'본점 주소',1,2,height=23,bg=grey,size=9,bold=True,color=mid)
    cellrow(8,_visit_cut(d['address'],80),3,8,size=9)
    cellrow(9,'업종 / 코드',1,2,height=23,bg=grey,size=9,bold=True,color=mid)
    cellrow(9,_visit_cut(d['industry'],80),3,8,size=9)
    cellrow(10,f"공시 기준: {_visit_cut(d['report_name'],75)}  |  DART {d['dart_code']}",height=18,size=8,color=mid)
    ws.row_dimensions[11].height=7
    section(12,'02  차입금 · 거래 금융기관 (감사보고서 원문 기반)')
    for a,b,label in [(1,2,'금융기관'),(3,4,'거래유형'),(5,7,'금액 표기(미검증)'),(8,8,'근거강도')]:
        cellrow(13,label,a,b,height=22,bg=light,size=8,bold=True,color=mid,align='center',indent=0)
    visible=d['banks'][:4]
    for k in range(4):
        r=14+k
        b=visible[k] if k<len(visible) else {}
        bg=white if k%2==0 else grey
        for a,e,v in [(1,2,_visit_cut(b.get('bank',''),19)),(3,4,_visit_cut(b.get('type',''),22)),
                      (5,7,_visit_cut(b.get('amount',''),32)),(8,8,b.get('strength',''))]:
            cellrow(r,v,a,e,height=21,size=9,bg=bg,align='center' if a==8 else 'left',indent=0 if a==8 else 1)
    remaining=max(0,len(d['banks'])-len(visible))
    if remaining:
        cellrow(18,f'※ 추가 금융기관 {remaining}곳은 기존 상세 엑셀/화면에서 확인',height=19,size=8,color=mid)
    else:
        cellrow(18,'※ 위 금액은 감사보고서 문맥에서 추출한 표기이며 차입금 잔액 합계로 확정하지 않습니다.',height=19,size=8,color=mid)
    summary_evidence=next((x['evidence'] for x in d['banks'] if x.get('evidence')), '')
    cellrow(19,'대표 공시 근거(발췌): '+_visit_cut(summary_evidence or '은행 관련 공시 근거 미확인',114),height=29,size=8,color=mid,bg=grey,wrap=True)
    ws.row_dimensions[20].height=7
    section(21,'03  퇴직연금 · 퇴직급여 공시 현황')
    cellrow(22,'운영 추정',1,2,height=22,bg=grey,size=9,bold=True,color=mid)
    cellrow(22,d['pension_type']+'  /  '+d['pension_confidence'],3,8,size=9,bold=True,color=teal)
    cellrow(23,'공시 단서',1,2,height=22,bg=grey,size=9,bold=True,color=mid)
    cellrow(23,f"DB {d['db_count']}건 · DC {d['dc_count']}건 · 사외적립자산 {d['asset_count']}건",3,8,size=9)
    cellrow(24,'금융기관',1,2,height=22,bg=grey,size=9,bold=True,color=mid)
    cellrow(24,_visit_cut(d['pension_inst'],76)+' (계약 여부 확인 필요)',3,8,size=9)
    cellrow(25,'※ 운영·계좌·연금사업자와 공시금액의 단위/기준일은 면담 및 원문으로 재확인',height=22,size=8,color=mid,bg=orange)
    ws.row_dimensions[26].height=8
    section(27,'04  현장 미팅 메모 / 후속 조치  — 직접 입력하거나 인쇄 후 작성')
    cellrow(28,'면담일시:                                  참석자 / 담당자:',height=24,size=9,color=dark,bg=light)
    cellrow(29,'● 거래은행 / 차입 조건 확인',height=20,size=9,bold=True,color=teal)
    for r in range(30,35):
        cellrow(r,'',height=21,bg=white)
        for col in range(1,9): ws.cell(r,col).border=Border(bottom=edge)
    cellrow(35,'● 퇴직연금 현황 / 제안 포인트',height=20,size=9,bold=True,color=teal)
    for r in range(36,41):
        cellrow(r,'',height=21,bg=white)
        for col in range(1,9): ws.cell(r,col).border=Border(bottom=edge)
    cellrow(41,'● 후속 일정 / 요청자료 / 기타',height=20,size=9,bold=True,color=teal)
    for r in range(42,48):
        cellrow(r,'',height=21,bg=white)
        for col in range(1,9): ws.cell(r,col).border=Border(bottom=edge)
    cellrow(48,'DART 감사보고서 기반 미확정 분석입니다. 면담을 통해 거래은행, 실제 차입금 및 연금사업자를 확인하세요.',height=20,size=8,color=mid,bg=light)

    ws.print_area='A1:H48'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='portrait'
    ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=1
    ws.print_options.horizontalCentered=True
    ws.page_margins=PageMargins(left=0.28,right=0.28,top=0.30,bottom=0.30,header=0.10,footer=0.10)
    ws.oddFooter.center.text='Corporate Finder  |  Internal review only  |  &P / &N'
    out=io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _visit_pdf_font():
    """외부 폰트 파일을 배포하지 않고 런타임의 한국어 폰트를 선택."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    opts=[
        ('/usr/share/fonts/truetype/nanum/NanumBarunGothic.ttf','/usr/share/fonts/truetype/nanum/NanumBarunGothicBold.ttf'),
        ('/usr/share/fonts/truetype/nanum/NanumGothic.ttf','/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf'),
        ('C:/Windows/Fonts/malgun.ttf','C:/Windows/Fonts/malgunbd.ttf'),
        ('C:/Windows/Fonts/gulim.ttc','C:/Windows/Fonts/gulim.ttc'),
    ]
    for reg,bold in opts:
        if os.path.isfile(reg):
            try:
                if 'CorpVisit-Regular' not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont('CorpVisit-Regular',reg))
                    pdfmetrics.registerFont(TTFont('CorpVisit-Bold',bold if os.path.isfile(bold) else reg))
                return 'CorpVisit-Regular','CorpVisit-Bold'
            except Exception:
                continue
    # 별도 폰트 배포가 불필요한 ReportLab 내장 한글 CID 폰트
    if 'HYGothic-Medium' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont('HYGothic-Medium'))
    return 'HYGothic-Medium','HYGothic-Medium'


def make_visit_report_pdf(result, external_meta=None):
    """A4 정확히 1페이지. 은행 근거는 요약하고 미팅 메모를 넓게 확보한다."""
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    d=_visit_report_context(result,external_meta)
    normal,bold=_visit_pdf_font()
    w,h=A4
    buf=io.BytesIO()
    pdf=canvas.Canvas(buf,pagesize=A4,pageCompression=1)
    pdf.setTitle(f"{d['name']} - 법인 현장 미팅 보고서")
    navy=HexColor('#18354F'); teal=HexColor('#008879'); ink=HexColor('#253849')
    mute=HexColor('#627485'); pale=HexColor('#EDF4F4'); grey=HexColor('#F2F5F7')
    stroke=HexColor('#D8E3E7'); cream=HexColor('#FFF7E8')
    margin=29; right=w-margin; inner_w=w-2*margin

    def box(top,height,color,rad=0,x=margin,width=inner_w,border=None):
        pdf.setFillColor(color)
        if border:pdf.setStrokeColor(border)
        if rad:
            pdf.roundRect(x,h-top-height,width,height,rad,fill=1,stroke=bool(border))
        else:
            pdf.rect(x,h-top-height,width,height,fill=1,stroke=bool(border))

    def text(x,top,value,size=9,color=ink,weight=False,maxwidth=None):
        v=_visit_text(value,'')
        f=bold if weight else normal
        pdf.setFont(f,size);pdf.setFillColor(color)
        if maxwidth:
            while len(v)>1 and pdfmetrics.stringWidth(v,f,size)>maxwidth:
                v=v[:-2].rstrip()+'…'
        pdf.drawString(x,h-top-size,v)

    def rule(y,x=margin+11,x2=right-11):
        pdf.setStrokeColor(stroke);pdf.setLineWidth(0.6)
        pdf.line(x,h-y,x2,h-y)

    def section(top,idx,label):
        text(margin,top,idx,size=10.0,color=teal,weight=True)
        text(margin+23,top-0.4,label,size=10.2,color=navy,weight=True)

    text(margin,18,'CORPORATE FINDER  /  FIELD BRIEF',size=9,color=teal,weight=True)
    box(37,43,navy,rad=4)
    text(margin+15,45,'법인 분석 요약 보고서',size=17,color=HexColor('#FFFFFF'),weight=True)
    text(right-139,62,f"작성 {d['generated']}",size=8,color=HexColor('#D3E5E5'))
    section(90,'01','기업 기본정보')
    box(111,65,grey,rad=5)
    text(margin+12,115,'회사명',size=8,color=mute)
    text(margin+70,113,d['name'],size=11,color=navy,weight=True,maxwidth=inner_w-95)
    rule(136)
    text(margin+12,139,'대표자',size=8,color=mute)
    text(margin+70,139,d['ceo'],size=8,maxwidth=112)
    text(margin+218,139,'전화',size=8,color=mute)
    text(margin+253,139,d['phone'],size=8,maxwidth=155)
    text(margin+12,158,'주소',size=8,color=mute)
    text(margin+70,158,d['address'],size=7.8,maxwidth=inner_w-90)
    text(margin,179,f"사업자 {d['bizno']}   |   법인등록번호 {d['jurino']}   |   업종 {d['industry']}",size=7.6,color=mute,maxwidth=inner_w)
    text(margin,193,f"공시: {d['report_name']}   (접수 {d['report_date']})",size=7.5,color=mute,maxwidth=inner_w)
    section(213,'02','차입금 · 거래 금융기관')
    box(235,120,HexColor('#FFFFFF'),rad=3,border=stroke)
    box(235,22,navy)
    header=[('금융기관',9),('거래유형',157),('원문 금액 표기',302),('강도',483)]
    for label,x in header: text(margin+x,240,label,size=8,color=HexColor('#FFFFFF'),weight=True)
    visible=d['banks'][:4]
    for i in range(4):
        rowtop=257+23*i
        if i%2:box(rowtop,23,grey,x=margin+1,width=inner_w-2)
        b=visible[i] if i<len(visible) else {}
        text(margin+9,rowtop+6,_visit_cut(b.get('bank',''),18),size=8,maxwidth=140)
        text(margin+157,rowtop+6,_visit_cut(b.get('type',''),19),size=7.7,maxwidth=142)
        text(margin+302,rowtop+6,_visit_cut(b.get('amount',''),24),size=7.7,maxwidth=170)
        text(margin+483,rowtop+6,_visit_cut(b.get('strength',''),6),size=7.7,maxwidth=42)
    extra=max(0,len(d['banks'])-len(visible))
    text(margin,359, ('추가 확인 금융기관 '+str(extra)+'곳: 상세분석 참조  |  ' if extra else '')+
         '금액은 원문 추출값이며 차입금 잔액 합계로 검증되지 않음',size=7.4,color=mute,maxwidth=inner_w)
    section(378,'03','퇴직연금 · 퇴직급여')
    box(400,92,pale,rad=5)
    text(margin+12,408,'운영분류',size=8,color=mute)
    text(margin+87,406,d['pension_type'],size=10,color=teal,weight=True,maxwidth=244)
    text(margin+347,409,d['pension_confidence'],size=7.5,color=mute,maxwidth=172)
    rule(429)
    text(margin+12,435,'공시 단서',size=8,color=mute)
    text(margin+87,435,f"DB {d['db_count']}건 / DC {d['dc_count']}건 / 사외적립자산 {d['asset_count']}건",size=8,maxwidth=425)
    text(margin+12,456,'기관 후보',size=8,color=mute)
    text(margin+87,456,d['pension_inst'],size=8,maxwidth=410)
    text(margin+12,476,'※ 사업자·가입 여부 및 실제 적립금은 미팅 때 확인',size=7.5,color=mute)
    section(510,'04','현장 미팅 메모  /  후속 조치')
    box(531,271,HexColor('#FFFFFF'),rad=4,border=stroke)
    text(margin+12,540,'면담일시:                                 참석자 / 담당자:',size=8.5,color=mute)
    rule(560)
    sections=[('거래은행 · 차입조건 확인',570,3),('퇴직연금 · 기존사업자',639,3),('후속 일정 · 요청사항 · 기타',708,4)]
    for title,top,count in sections:
        text(margin+13,top,title,size=8.7,color=teal,weight=True)
        for i in range(count):rule(top+21+17*i,margin+13,right-13)
    rule(812,x=margin,x2=right)
    text(margin,816,'DART 감사보고서 기반 요약 자료 / 거래관계·연금운영·금액은 원문 및 면담으로 검증',size=7.0,color=mute,maxwidth=inner_w)
    pdf.showPage();pdf.save()
    return buf.getvalue()


def _visit_report_preview_html(d):
    """미리보기는 PDF와 동일한 공시 내용만 간결하게 노출한다."""
    esc=lambda v:html.escape(str(v or ''))
    trs=''.join(
        '<tr><td>'+esc(b['bank'])+'</td><td>'+esc(_visit_cut(b['type'],26))+'</td>'+
        '<td>'+esc(_visit_cut(b['amount'],30))+'</td><td>'+esc(b['strength'])+'</td></tr>'
        for b in d['banks'][:4]
    ) or '<tr><td colspan="4">보고서에서 금융기관 근거를 확인하지 못했습니다.</td></tr>'
    more=(f'<small>외 {len(d["banks"])-4}개 금융기관: 상세 분석 참조</small>' if len(d['banks'])>4 else '')
    return f"""<!doctype html><html lang="ko"><meta charset="utf-8"><style>
    *{{box-sizing:border-box}}body{{margin:0;padding:14px;background:#eaf0f1;font-family:"Malgun Gothic","Noto Sans KR",sans-serif;color:#253849}}
    article{{max-width:710px;margin:auto;background:white;min-height:870px;padding:26px;box-shadow:0 2px 12px #82909830}}
    header{{background:#18354f;color:white;padding:18px 20px;border-radius:6px}}header strong{{font-size:23px}}header div{{font-size:12px;margin-top:5px;opacity:.87}}
    h3{{font-size:14px;margin:18px 0 7px;color:#008879}}.basics{{display:grid;grid-template-columns:1fr 1fr;gap:5px 13px;background:#f2f5f7;border-radius:5px;padding:12px;font-size:11px}}
    .wide{{grid-column:span 2}}table{{border-collapse:collapse;width:100%;font-size:11px}}th{{background:#18354f;color:white}}td,th{{text-align:left;padding:9px;border-bottom:1px solid #d8e3e7}}tr:nth-child(even){{background:#f2f5f7}}
    .pension{{background:#edf4f4;padding:12px;font-size:11px;line-height:1.8;border-radius:5px}}
    .memo{{height:265px;border:1px solid #d8e3e7;border-radius:5px;padding:10px 14px}}.memo b{{display:block;color:#008879;margin:8px 0}}.lines{{height:45px;background:repeating-linear-gradient(to bottom,transparent 0px,transparent 19px,#d8e3e7 20px,#d8e3e7 21px)}}
    footer{{font-size:10px;color:#627485;margin-top:11px}}
    </style><body><article><header><strong>법인 분석 요약 보고서</strong><div>{esc(d['name'])} · {esc(d['generated'])}</div></header>
    <h3>01 기업 기본정보</h3><div class="basics"><div>대표자 : {esc(d['ceo'])}</div><div>전화번호 : {esc(d['phone'])}</div>
    <div>사업자번호 : {esc(d['bizno'])}</div><div>법인등록번호 : {esc(d['jurino'])}</div>
    <div class="wide">본점 주소 : {esc(d['address'])}</div><div class="wide">감사보고서 : {esc(d['report_name'])} / {esc(d['report_date'])}</div></div>
    <h3>02 차입금 · 거래 금융기관</h3><table><tr><th>금융기관</th><th>거래유형</th><th>금액 표기(미검증)</th><th>강도</th></tr>{trs}</table>{more}
    <h3>03 퇴직연금 · 퇴직급여</h3><div class="pension"><b>{esc(d['pension_type'])}</b> ({esc(d['pension_confidence'])})<div>DB {d['db_count']}건 · DC {d['dc_count']}건 · 사외적립자산 {d['asset_count']}건</div><div>연금 금융기관 후보: {esc(d['pension_inst'])}</div><div>※ 실제 사업자·계약 및 적립금은 회사 확인 필요</div></div>
    <h3>04 현장 미팅 메모 / 후속 조치</h3><div class="memo"><b>거래은행 · 차입조건</b><div class="lines"></div><b>퇴직연금 · 기존사업자</b><div class="lines"></div><b>후속 일정 · 요청사항 · 기타</b><div class="lines"></div></div>
    <footer>공시 기반의 내부 검토자료이며 실제 거래은행·차입금 잔액·퇴직연금 사업자를 확정하지 않습니다.</footer></article></body></html>"""


def render_visit_report_exports(result, widget_key, external_meta=None):
    """법인명·지역 검색 공통. 기존 상세 엑셀은 호출부에서 별도로 유지."""
    d=_visit_report_context(result,external_meta)
    st.markdown('### 📄 A4 현장 미팅용 법인 분석 보고서')
    st.caption('기업 기본정보 · 금융기관/차입 관련 공시 · 퇴직연금 현황 / 하단은 현장 메모 공간. 인쇄는 A4 세로 1장입니다.')
    with st.expander('👁️ A4 보고서 미리보기',expanded=False):
        import streamlit.components.v1 as components
        components.html(_visit_report_preview_html(d),height=940,scrolling=True)
    ca,cb=st.columns(2)
    with ca:
        try:
            pdf_bytes=make_visit_report_pdf(result,external_meta)
            st.download_button('📄 A4 보고서 PDF 다운로드', data=pdf_bytes,
                               file_name=f"{_visit_filename(d['name'])}_현장미팅보고서.pdf",
                               mime='application/pdf',key=widget_key+'_pdf',use_container_width=True)
        except ImportError:
            st.warning('PDF 생성에는 reportlab이 필요합니다. requirements.txt에 reportlab>=4.2를 추가해 주세요.')
        except Exception as exc:
            st.warning(f'PDF 생성 실패: {type(exc).__name__} (기존 분석은 유지됩니다)')
    with cb:
        try:
            excel_bytes=make_visit_report_excel(result,external_meta)
            st.download_button('📊 A4 편집용 엑셀 다운로드',data=excel_bytes,
                               file_name=f"{_visit_filename(d['name'])}_현장미팅보고서.xlsx",
                               mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                               key=widget_key+'_a4excel',use_container_width=True)
        except Exception as exc:
            st.warning(f'A4 엑셀 생성 실패: {type(exc).__name__} (기존 상세 엑셀은 유지됩니다)')
    st.caption('※ 금융기관·금액·퇴직연금 정보는 감사보고서 추출 단서입니다. 확정 거래/잔액/연금사업자 정보로 사용하지 마세요.')


# =========================================================
# =========================================================
# 신규: 퇴직급여 사외적립 제도 대응 타깃 발굴 (2026-10-30)
# 기존 거래은행/DART/SQLite/Supabase 분석과 독립적으로 동작합니다.
# 근로자 수는 감사보고서 키워드만으로 추정하지 않습니다.
# =========================================================
PENSION_POLICY_SOURCE = 'https://www.moel.go.kr/news/enews/report/enewsView.do?news_seq=20160'
PENSION_LEAD_STAGES = ['미접촉', '접촉 예정', '전화 완료', '미팅 예정', '제안서 전달', '수주', '보류']


def _pension_target_init():
    """기존 테이블은 변경하지 않고 영업 메모용 보조 테이블만 추가."""
    with get_db_connection() as conn:
        conn.execute('''CREATE TABLE IF NOT EXISTS pension_target_meta (
            corp_code TEXT PRIMARY KEY,
            employee_count INTEGER,
            contact_stage TEXT NOT NULL DEFAULT '미접촉',
            contact_note TEXT NOT NULL DEFAULT '',
            updated_at TEXT NOT NULL DEFAULT ''
        )''')


def _pension_employee_init():
    with get_db_connection() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS pension_employee_dart (
            corp_code TEXT PRIMARY KEY,
            employee_count INTEGER NOT NULL,
            report_year TEXT NOT NULL,
            report_code TEXT NOT NULL,
            report_date TEXT NOT NULL DEFAULT '',
            report_url TEXT NOT NULL DEFAULT '',
            checked_at TEXT NOT NULL
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS pension_employee_check (
            corp_code TEXT PRIMARY KEY,
            status TEXT NOT NULL,
            detail TEXT NOT NULL DEFAULT '',
            checked_at TEXT NOT NULL
        )""")


def _pension_int(value):
    if value is None or isinstance(value, bool):
        return None
    s = str(value).strip().replace(',', '').replace('명', '').replace(' ', '')
    if not s or s.lower() in ('nan', 'none', '미확인', '-', 'null'):
        return None
    if not re.fullmatch(r'\d+(?:\.0+)?', s):
        return None
    n = int(float(s))
    return n if 1 <= n <= 1000000 else None


# DART 정기보고서 직원현황 API: 감사보고서(document.xml)와는 별도 출처입니다.
# 원문에 직원 수가 없으면 추정하지 않으며 사업보고서 미제출 법인은 자료가 없을 수 있습니다.
def _pension_employee_periods(now=None):
    now = now or date.today()
    y = now.year
    periods = []
    if now.month >= 11:
        periods.append((str(y), '11014'))  # 3분기 보고서
    if now.month >= 9:
        periods.append((str(y), '11012'))  # 반기 보고서
    if now.month >= 5:
        periods.append((str(y), '11013'))  # 1분기 보고서
    # 최신 사업보고서 및 전년도 보고서; 장기 자료의 현직원수 오인 방지
    periods.append((str(y-1), '11011'))
    return periods[:4]


def _pension_employee_count_from_rows(items):
    """DART sm(합계)만 활용. 부문/성별 소계와 상세를 중복 합산하지 않는다."""
    valid=[]
    for row in items or []:
        if not isinstance(row, dict):
            continue
        n=_pension_int(row.get('sm'))
        if n is None:
            continue
        group=re.sub(r'\s+', '', str(row.get('fo_bbm') or ''))
        sex=re.sub(r'\s+', '', str(row.get('sexdstn') or ''))
        valid.append((n, group, sex))
    if not valid:
        return None
    overall = [r for r in valid if re.fullmatch(r'(?:총계|합계|전체|계|전사)', r[1])]
    if overall:
        all_sex = [r[0] for r in overall if r[2] in ('계','합계','전체','총계','')]
        if all_sex:
            return max(all_sex)
        return sum(n for n,g,sex in overall if sex in ('남','여','기타')) or None
    # 부문별 행: 동일 부문에서 성별 '계'와 남/여가 동시에 있으면 계만 사용
    groups={}
    for n,g,sex in valid:
        groups.setdefault(g,[]).append((n,sex))
    total=0
    for grp in groups.values():
        subs=[n for n,sex in grp if sex in ('계','합계','전체','총계','')]
        if subs:
            total += max(subs)
        else:
            total += sum(n for n,sex in grp if sex in ('남','여','기타'))
    return total if total>0 else None


def _pension_employee_lookup(api_key, corp_code, now=None):
    """명시적 사용/분석 실행 시만 호출. 정상 013은 다음 보고서를 탐색."""
    if not api_key or not re.fullmatch(r'\d{8}', str(corp_code)):
        raise ValueError('DART API 키와 8자리 법인번호가 필요합니다.')
    empty_periods=[]
    for yr,code in _pension_employee_periods(now):
        result=request_json(f'{DART_BASE}/empSttus.json', {
            'crtfc_key':api_key, 'corp_code':str(corp_code),
            'bsns_year':yr,'reprt_code':code},timeout=12)
        if str(result.get('status','000'))=='013':
            empty_periods.append(f'{yr}/{code}')
            continue
        count=_pension_employee_count_from_rows(result.get('list') or [])
        if count is not None:
            rows=result.get('list') or []
            dates=[str(r.get('stlm_dt','')) for r in rows if isinstance(r,dict) and r.get('stlm_dt')]
            receipts=[str(r.get('rcept_no','')) for r in rows if isinstance(r,dict)]
            receipt=next((x for x in receipts if re.fullmatch(r'\d{14}',x)), '')
            return {'employee_count':count,'report_year':yr,'report_code':code,
                    'report_date':max(dates) if dates else '',
                    'report_url':f'https://dart.fss.or.kr/dsaf001/main.do?rcpNo={receipt}' if receipt else ''}
        empty_periods.append(f'{yr}/{code}:인원 숫자 없음')
    _pension_employee_log(corp_code,'미공개','최신 정기보고서 직원현황 없음')
    return None


def _pension_employee_save(corp_code, data):
    _pension_employee_init()
    if not re.fullmatch(r'\d{8}',str(corp_code)):
        raise ValueError('법인번호 형식이 올바르지 않습니다.')
    count=_pension_int(data.get('employee_count'))
    if count is None:
        raise ValueError('공시에서 유효한 인원수를 확인하지 못했습니다.')
    with get_db_connection() as conn:
        conn.execute("""INSERT INTO pension_employee_dart
          (corp_code,employee_count,report_year,report_code,report_date,report_url,checked_at)
          VALUES (?,?,?,?,?,?,?) ON CONFLICT(corp_code) DO UPDATE SET
          employee_count=excluded.employee_count,report_year=excluded.report_year,
          report_code=excluded.report_code,report_date=excluded.report_date,
          report_url=excluded.report_url,checked_at=excluded.checked_at""",
          (str(corp_code),count,str(data['report_year']),str(data['report_code']),
           str(data.get('report_date','')),str(data.get('report_url','')),
           datetime.now().isoformat(timespec='seconds')))
    _pension_employee_log(corp_code,'확인',f'{count}명 / {data["report_year"]}년')


def _pension_employee_saved(corp_code):
    _pension_employee_init()
    with get_db_connection() as conn:
        item=conn.execute('SELECT * FROM pension_employee_dart WHERE corp_code=?',(str(corp_code),)).fetchone()
    return dict(item) if item is not None else None


def _pension_employee_log(corp_code, status, detail=''):
    _pension_employee_init()
    if not re.fullmatch(r'\d{8}',str(corp_code)):
        return
    with get_db_connection() as conn:
        conn.execute("""INSERT INTO pension_employee_check (corp_code,status,detail,checked_at)
            VALUES (?,?,?,?) ON CONFLICT(corp_code) DO UPDATE SET
            status=excluded.status,detail=excluded.detail,checked_at=excluded.checked_at""",
            (str(corp_code),str(status)[:30],str(detail or '')[:180],datetime.now().isoformat(timespec='seconds')))


def _pension_employee_recent_checks():
    _pension_employee_init()
    with get_db_connection() as conn:
        rows=conn.execute('SELECT * FROM pension_employee_check').fetchall()
    return {str(r['corp_code']):dict(r) for r in rows}


def _pension_employee_effective(corp_code, fallback=None):
    """직접 확인 인원 > 외부 기업목록 인원 > DART 보고서 참고 인원."""
    _pension_target_init()
    _pension_employee_init()
    with get_db_connection() as conn:
        manual=conn.execute('SELECT employee_count FROM pension_target_meta WHERE corp_code=?',(str(corp_code),)).fetchone()
    n=_pension_int(manual['employee_count']) if manual is not None else None
    if n is not None:
        return n,'담당자 확인/수동 입력', None
    ext=_pension_int(fallback)
    if ext is not None:
        return ext,'외부 기업목록',None
    saved=_pension_employee_saved(corp_code)
    if saved:
        return saved['employee_count'],f'DART 정기보고서 {saved["report_year"]}년 참고',saved
    return None,'확인 불가',None


def _render_pension_employee_section(corp_code, api_key, enabled, external_count=None, key='pension'):
    """기존 개별 법인 분석 화면에도 직원 수와 시행연도, 원문을 항상 표시한다."""
    st.markdown('#### 👥 종업원 수 · 의무화 적용시기 검토')
    number, source, dart_row = _pension_employee_effective(corp_code,external_count)
    a,b,c=st.columns([1,1.2,2])
    a.metric('직원 수',f'{number:,}명' if number is not None else '미확인')
    b.metric('사외적립 의무화',_pension_policy_year(number))
    c.caption(f'**인원 출처:** {source}')
    if dart_row:
        c.caption(f'공시 기준: {dart_row["report_date"] or dart_row["report_year"]} · 조회 저장: {dart_row["checked_at"][:10]}')
        if dart_row.get('report_url'):
            c.link_button('직원현황 공시 원문',dart_row['report_url'])
    st.caption('DART 직원현황은 사업·반기·분기보고서상의 참고 인원이며 법정 상시근로자 수나 퇴직연금 가입 인원과 다를 수 있습니다.')
    recent=_pension_employee_recent_checks().get(str(corp_code))
    if recent and not dart_row and number is None:
        st.caption(f'DART 조회기록: {recent["status"]} · {recent["checked_at"][:10]} (새로 조회하려면 아래 버튼)')
    if enabled:
        if st.button('🔎 DART 직원현황 조회·갱신',key=key+'_emp_refresh'):
            try:
                with st.spinner('최신 정기보고서의 직원현황을 확인하고 있습니다...'):
                    record=_pension_employee_lookup(api_key,str(corp_code))
                if record:
                    _pension_employee_save(corp_code,record)
                    if not _db_cloud_push():
                        st.warning('로컬 저장 완료, Supabase 백업 상태는 별도 확인해 주세요.')
                    st.success(f'DART 공시 기준 {record["employee_count"]:,}명 저장')
                    st.rerun()
                else:
                    st.info('DART 정기보고서 직원현황 자료가 없습니다. 감사보고서만 제출하는 법인 등은 직원 수를 조회하지 못할 수 있습니다. 타깃 탭에서 수동/엑셀 등록이 가능합니다.')
            except Exception as exc:
                _pension_employee_log(corp_code,'조회 오류',type(exc).__name__)
                st.warning(f'DART 직원현황 조회 실패: {type(exc).__name__}: {str(exc)[:170]}')
    else:
        st.caption('오프라인 모드 또는 DART API 키 미설정: 저장된 정보만 표시됩니다.')


def _pension_policy_year(count):
    if count is None:
        return '인원 확인 필요'
    if count >= 100:
        return '2028년(추진안)'
    if count >= 10:
        return '2029년(추진안)'
    return '2030년(추진안)'


def _pension_target_classify(result):
    """미도입 확정/추정 비율을 생성하지 않음. 공시 근거 단계만 표시."""
    if not result or result.get('error') or not result.get('report'):
        return ('미분석', '감사보고서 분석 필요', 4, '감사보고서 미분석·원문 점검')
    pen = result.get('pension') or {}
    op = str(pen.get('운영분류') or '확인 불가')
    if op == '미운용 확인':
        return ('미도입 후보', '명시적 미운용 문구 · 원문 재확인 필수', 1, '공시일·적용대상·현재 도입 여부 확인')
    if op in ('DB형 운영', 'DC형 운영', 'DB+DC 혼합'):
        return ('기존 운영', op + ' · 사업자/운용 점검 대상', 5, '퇴직연금사업자·적립수준·운용 현황 확인')
    if str(pen.get('판별신뢰')) == '상충 근거':
        return ('상충 근거', '운영·미운영 근거가 함께 발견됨', 3, '회사 전체 기준 및 DB/DC 병행 여부 확인')
    evidence = bool(pen.get('일반근거') or pen.get('DB근거') or pen.get('DC근거'))
    if evidence:
        return ('가입 여부 확인', '퇴직급여 관련 문구만 존재', 2, '실제 가입 여부·퇴직급여제도 확인')
    return ('자료 부족', '퇴직연금 운영 관련 직접 근거 없음', 3, '감사보고서 주석 및 담당자 추가 확인')


def _pension_target_load(corp_master, external_df=None):
    """저장된 감사보고서 분석과 외부 법인목록을 결합. 신규 DART 호출 없음."""
    # 초기화 순서가 달라지는 Cloud 재시작·구형 DB 복원도 안전하게 처리
    _pension_target_init()
    _pension_employee_init()
    with get_db_connection() as conn:
        analyses = conn.execute('''SELECT a.corp_code, a.result_json, a.analyzed_at,
                p.corp_name AS profile_name, p.adres AS profile_addr, p.phn_no AS profile_phone
                FROM company_analysis a
                LEFT JOIN company_profile p ON a.corp_code=p.corp_code''').fetchall()
        extra = conn.execute('SELECT corp_code, employee_count, contact_stage, contact_note, updated_at FROM pension_target_meta').fetchall()
        dart_employee = conn.execute('SELECT * FROM pension_employee_dart').fetchall()
    meta = {str(x['corp_code']): dict(x) for x in extra}
    dart_meta = {str(x['corp_code']): dict(x) for x in dart_employee}
    master = {}
    norm_to_codes = {}
    if corp_master is not None and not corp_master.empty:
        for row in corp_master[['corp_code', 'corp_name', 'norm_name']].fillna('').to_dict('records'):
            code = str(row['corp_code'])
            master[code] = str(row['corp_name'])
            norm = str(row['norm_name']) or normalize_company_name(str(row['corp_name']))
            norm_to_codes.setdefault(norm, set()).add(code)

    analyses_by_code = {}
    for x in analyses:
        code = str(x['corp_code'])
        try:
            result = json.loads(x['result_json'])
            if not isinstance(result, dict):
                continue
        except (ValueError, TypeError):
            continue
        analyses_by_code[code] = (result, str(x['analyzed_at'] or ''), dict(x))

    external_by_code = {}
    unmatched = {}
    if external_df is not None and isinstance(external_df, pd.DataFrame) and not external_df.empty:
        for row in external_df.fillna('').to_dict('records'):
            name = str(row.get('법인명') or '').strip()
            if not name:
                continue
            norm = str(row.get('norm_name') or normalize_company_name(name))
            codes = norm_to_codes.get(norm, set())
            if len(codes) == 1:
                code = next(iter(codes))
                if code not in external_by_code:
                    external_by_code[code] = row
            elif norm not in unmatched:
                unmatched[norm] = (row, '동명법인 확인 필요' if codes else 'DART 매칭 필요')

    all_codes = set(analyses_by_code) | set(external_by_code)
    rows = []
    for code in all_codes:
        record = analyses_by_code.get(code)
        result, analyzed_at, sqlrow = record if record else (None, '', {})
        company = (result or {}).get('company') or {}
        report = (result or {}).get('report') or {}
        pension = (result or {}).get('pension') or {}
        external = external_by_code.get(code, {})
        saved = meta.get(code, {})
        category, explanation, rank, suggested = _pension_target_classify(result)
        employees = _pension_int(saved.get('employee_count'))
        employee_source = '담당자 확인/수동' if employees is not None else ''
        if employees is None:
            employees = _pension_int(external.get('종업원수'))
            if employees is not None:
                employee_source = '외부 기업목록'
        dart_observation = dart_meta.get(code, {})
        if employees is None and dart_observation:
            employees = _pension_int(dart_observation.get('employee_count'))
            if employees is not None:
                employee_source = 'DART 정기보고서(참고)'
        addr = str(company.get('주소') or sqlrow.get('profile_addr') or external.get('주소') or '')
        rows.append({
            '법인명': str(company.get('법인명') or master.get(code) or sqlrow.get('profile_name') or external.get('법인명') or ''),
            'DART고유번호': code,
            '주소': addr,
            '전화번호': str(company.get('전화번호') or sqlrow.get('profile_phone') or external.get('전화번호') or ''),
            '종업원수': employees,
            '인원출처': employee_source or '미확인',
            '인원공시기준': str(dart_observation.get('report_date') or dart_observation.get('report_year') or ''),
            '시행예정': _pension_policy_year(employees),
            '대상유형': category,
            '판별근거': explanation,
            '영업제안': suggested,
            '연금운영분류': str(pension.get('운영분류') or '미확인'),
            '퇴직급여단서수': len(pension.get('일반근거') or []),
            '사외적립원문수': len(((pension.get('금액관련근거') or {}).get('사외적립자산') or [])),
            '접수일': str(report.get('접수일') or ''),
            '분석저장일': analyzed_at[:10],
            'DART원문': str(report.get('DART링크') or ''),
            '영업상태': str(saved.get('contact_stage') or '미접촉'),
            '상담메모': str(saved.get('contact_note') or ''),
            '출처': ('저장 분석 + 외부기업목록' if code in analyses_by_code and code in external_by_code
                   else '저장 분석' if code in analyses_by_code else '외부기업목록'),
            '_rank': rank,
        })
    for norm, (ext, matching) in unmatched.items():
        rows.append({
            '법인명': str(ext.get('법인명') or ''), 'DART고유번호': '',
            '주소': str(ext.get('주소') or ext.get('검색주소') or ''),
            '전화번호': str(ext.get('전화번호') or ''),
            '종업원수': _pension_int(ext.get('종업원수')),
            '인원출처': '외부 기업목록' if _pension_int(ext.get('종업원수')) is not None else '미확인',
            '인원공시기준': '',
            '시행예정': _pension_policy_year(_pension_int(ext.get('종업원수'))),
            '대상유형': '미분석', '판별근거': matching,
            '영업제안': '법인명·DART고유번호 확인 후 신규 분석',
            '연금운영분류': '미확인', '퇴직급여단서수': 0,
            '사외적립원문수': 0, '접수일': '', '분석저장일': '', 'DART원문': '',
            '영업상태': '미접촉', '상담메모': '', '출처': '외부기업목록', '_rank': 4,
        })
    columns = ['법인명', 'DART고유번호', '주소', '전화번호', '종업원수', '인원출처', '인원공시기준', '시행예정',
               '대상유형', '판별근거', '영업제안', '연금운영분류', '퇴직급여단서수',
               '사외적립원문수', '접수일', '분석저장일', 'DART원문', '영업상태',
               '상담메모', '출처', '_rank']
    df = pd.DataFrame(rows, columns=columns)
    if not df.empty:
        df = df.sort_values(['_rank', '법인명'], kind='stable').reset_index(drop=True)
    return df


def _pension_target_save(code, *, headcount=None, stage='미접촉', note=''):
    if not re.fullmatch(r'\d{8}', str(code)):
        raise ValueError('DART 8자리 코드가 확인된 법인만 영업메모를 저장할 수 있습니다.')
    if stage not in PENSION_LEAD_STAGES:
        raise ValueError('영업상태 값이 올바르지 않습니다.')
    if headcount is not None and _pension_int(headcount) is None:
        raise ValueError('종업원 수는 1명 이상의 정수로 입력해 주세요.')
    with get_db_connection() as conn:
        conn.execute('''INSERT INTO pension_target_meta
             (corp_code, employee_count, contact_stage, contact_note, updated_at)
             VALUES (?,?,?,?,?) ON CONFLICT(corp_code) DO UPDATE SET
             employee_count=excluded.employee_count,
             contact_stage=excluded.contact_stage,
             contact_note=excluded.contact_note,
             updated_at=excluded.updated_at''',
             (code, headcount, stage, str(note or '')[:2000], datetime.now().isoformat(timespec='seconds')))


def _pension_target_import_headcount(upload, corp_master):
    """일괄 입력은 8자리 DART 코드 또는 정확히 일치하는 유일한 법인명만 매칭."""
    name = str(upload.name).lower()
    raw = upload.getvalue()
    if name.endswith('.xlsx'):
        src = pd.read_excel(io.BytesIO(raw), dtype=str).fillna('')
    elif name.endswith('.csv'):
        last_error = None
        for encoding in ('utf-8-sig', 'cp949', 'euc-kr'):
            try:
                src = pd.read_csv(io.BytesIO(raw), dtype=str, encoding=encoding).fillna('')
                break
            except Exception as exc:
                last_error = exc
        else:
            raise ValueError(f'CSV 읽기 실패: {last_error}')
    else:
        raise ValueError('xlsx 또는 csv만 지원합니다.')
    col_name = _find_external_column(src, ['법인명', '회사명', '업체명', '기업명'])
    col_code = _find_external_column(src, ['DART고유번호', '고유번호', 'corp_code'])
    col_count = _find_external_column(src, ['종업원수', '직원수', '근로자수', '상시근로자수', '임직원수', '인원'])
    if not col_count or (not col_name and not col_code):
        raise ValueError('법인명 또는 DART고유번호 열과 종업원수(직원수) 열이 필요합니다.')
    by_name = {}
    valid_codes = set()
    for item in corp_master[['corp_code', 'corp_name']].fillna('').to_dict('records'):
        code = str(item['corp_code'])
        valid_codes.add(code)
        by_name.setdefault(normalize_company_name(str(item['corp_name'])), set()).add(code)
    to_save = {}
    missed = 0
    for item in src.to_dict('records'):
        employees = _pension_int(item.get(col_count))
        if employees is None:
            missed += 1
            continue
        code = str(item.get(col_code) or '').strip() if col_code else ''
        if re.fullmatch(r'\d{1,8}', code):
            code = code.zfill(8)
        if code not in valid_codes:
            codes = by_name.get(normalize_company_name(str(item.get(col_name) or '')), set()) if col_name else set()
            code = next(iter(codes)) if len(codes) == 1 else ''
        if not code:
            missed += 1
            continue
        to_save[code] = employees
    return to_save, missed, len(src)


def _pension_target_save_headcount(mapping):
    now = datetime.now().isoformat(timespec='seconds')
    with get_db_connection() as conn:
        for code, count in mapping.items():
            conn.execute('''INSERT INTO pension_target_meta
               (corp_code, employee_count, contact_stage, contact_note, updated_at)
               VALUES (?,?,'미접촉','',?) ON CONFLICT(corp_code) DO UPDATE SET
               employee_count=excluded.employee_count,
               updated_at=excluded.updated_at''', (code, count, now))


def _pension_target_excel(df):
    """화면 필터에 해당하는 목록과 작성 주의사항을 XLSX로 내보내기."""
    buf = io.BytesIO()
    out = df.drop(columns=['_rank'], errors='ignore').copy()
    # 외부기업목록의 셀 문자열이 Excel 수식으로 실행되지 않도록 보호.
    for col in out.select_dtypes(include=['object']).columns:
        out[col] = out[col].map(lambda v: ("'" + v) if isinstance(v, str) and v.startswith(('=', '+', '@')) else v)
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        out.to_excel(writer, sheet_name='퇴직연금_타깃', index=False)
        pd.DataFrame([
            ('자료 기준', '저장된 감사보고서 분석 + DART 정기보고서 직원현황 + 검증된 기업목록/수동 인원'),
            ('직원 수 출처', 'DART 직원현황은 참고 수치로 상시근로자수 또는 퇴직연금 가입자 수를 의미하지 않음'),
            ('판정 주의', '미도입 후보는 명시적 공시 문구에 따른 검토 대상이며 현재 미가입 확정이 아님'),
            ('미분석', '기존 DART 공시 분석 결과가 없는 법인. 미도입 추정 금지'),
            ('2028년 추진안', '100인 이상 사업장 (국회 법 개정 전)'),
            ('2029년 추진안', '10~99인 사업장 (국회 법 개정 전)'),
            ('2030년 추진안', '10인 미만 사업장 (국회 법 개정 전)'),
            ('최초 발표', '고용노동부 2026-10-29'),
            ('정책 원문', PENSION_POLICY_SOURCE),
            ('개별 조회', 'DART 감사보고서 공시 접수일·가입 범위와 회사 현재 가입 여부 재확인 필요'),
        ], columns=['구분', '설명']).to_excel(writer, sheet_name='판정_가이드', index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = 'C2' if ws.title == '퇴직연금_타깃' else 'A2'
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 25
            from openpyxl.styles import Font, PatternFill, Alignment
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='0B5045')
                cell.font = Font(color='FFFFFF', bold=True)
                cell.alignment = Alignment(vertical='center')
            for col in ws.columns:
                letter = col[0].column_letter
                width = min(42, max(12, max(len(str(cell.value or '')) for cell in list(col)[:60]) + 2))
                ws.column_dimensions[letter].width = width
    buf.seek(0)
    return buf.getvalue()


def render_pension_target_tab(corp_master):
    st.subheader('🎯 퇴직연금 의무화 대응 · 신규 법인 타깃')
    st.caption('기존 저장 결과에서 자동 선별 · 신규 DART 호출 없음 · 기존 법인 DB/분석 결과를 수정하지 않음')
    st.info('2026.10.29 정부 발표 추진안: 2028년 100인 이상 → 2029년 10~99인 → 2030년 10인 미만. ' 
            '법 개정 완료 전이므로 확정된 시행 의무로 안내하지 마세요. 시행 이후 근속분부터 적용 예정.')
    st.link_button('고용노동부 공식 발표 확인', PENSION_POLICY_SOURCE)

    ext = st.session_state.get('external_company_df')
    df = _pension_target_load(corp_master, ext)
    if df.empty:
        st.warning('저장된 감사보고서 분석과 외부 기업목록이 없습니다. 먼저 법인을 분석하거나 지역 검색에서 외부 기업목록을 등록하세요.')
        return

    total_saved = int(df['출처'].str.contains('저장 분석', regex=False).sum())
    cols = st.columns(4)
    cols[0].metric('총 검토 법인', f'{len(df):,}개')
    cols[1].metric('명시적 미운용 단서', f'{(df["대상유형"] == "미도입 후보").sum():,}개')
    cols[2].metric('가입 여부 확인', f'{(df["대상유형"] == "가입 여부 확인").sum():,}개')
    cols[3].metric('기존 분석 결과', f'{total_saved:,}개')

    st.markdown('#### 👥 직원 수 자동조회 · DART 정기보고서')
    report_count = int((df['인원출처'] == 'DART 정기보고서(참고)').sum())
    known_count = int(df['종업원수'].notna().sum())
    st.caption(f'현재 인원 확인 {known_count:,}/{len(df):,}개 법인 · DART 자동수집 {report_count:,}개. ' 
               '이전 버전은 직원 수를 자동 조회하지 않았습니다. 이 버튼은 실제 DART API로 조회합니다.')
    scope = st.selectbox('자동조회 범위', ['현재 타깃 중 미확인 법인', '전체 저장분석 중 미확인 법인'],
                          key='pension_dart_employee_scope')
    candidates = df[(df['DART고유번호'].str.fullmatch(r'\d{8}',na=False)) & (df['종업원수'].isna())]
    if scope.startswith('현재 타깃'):
        candidates = candidates[candidates['대상유형'].isin(['미도입 후보','가입 여부 확인','자료 부족','상충 근거'])]
    checks = _pension_employee_recent_checks()
    cutoff_missing = (datetime.now()-timedelta(days=30)).isoformat(timespec='seconds')
    cutoff_errors = (datetime.now()-timedelta(days=1)).isoformat(timespec='seconds')
    recent_codes={code for code,value in checks.items() if
                  str(value.get('checked_at','')) >=
                  (cutoff_missing if value.get('status') == '미공개' else cutoff_errors)}
    candidates=candidates[~candidates['DART고유번호'].isin(recent_codes)].head(10)
    already_checked=sum(1 for code in recent_codes if checks[code].get('status')!='확인')
    st.caption(f'최근 30일 미공개·지난 24시간 오류로 중복 조회 제외 {already_checked}개 (개별 법인에서는 강제 갱신 가능)')
    previous=st.session_state.pop('_pension_last_batch_result',None)
    if previous:
        st.success(previous)
    if st.button(f'🌐 DART 직원현황 자동조회 (최대 {len(candidates)}개씩)',
                 key='pension_employee_batch',disabled=(not dart_enabled or candidates.empty)):
        successes=[]; no_data=[]; failures=[]
        progress=st.progress(0,text='직원 현황 조회 중')
        for idx,row in enumerate(candidates.to_dict('records')):
            code=str(row['DART고유번호'])
            try:
                res=_pension_employee_lookup(api_key,code)
                if res:
                    _pension_employee_save(code,res)
                    successes.append(f'{row["법인명"]} {res["employee_count"]:,}명')
                else:
                    no_data.append(row['법인명'])
            except Exception as exc:
                failures.append(f'{row["법인명"]}: {type(exc).__name__}')
                _pension_employee_log(code,'조회 오류',type(exc).__name__)
                # 네트워크 장애 시 나머지 법인 연쇄 재시도 방지
                if _dart_cooldown_remaining():
                    break
            progress.progress((idx+1)/len(candidates),text=f'조회 {idx+1}/{len(candidates)}')
        progress.empty()
        if successes or no_data or failures:
            pushed=_db_cloud_push()
            if successes:
                st.session_state['_pension_last_batch_result']=f'직원 수 확인 {len(successes)}개: '+', '.join(successes[:5])
            if not pushed:
                st.warning('로컬 저장 성공, Supabase 영구저장은 추가 확인 필요')
        if no_data:
            st.info('API에 직원현황 미공개 ('+str(len(no_data))+'개): '+', '.join(no_data[:7]))
        if failures:
            st.warning('조회 오류 ('+str(len(failures))+'개): '+', '.join(failures[:4]))
        if successes:
            st.rerun()
    if not dart_enabled:
        st.caption('DART API 키를 설정하고 오프라인 모드를 해제하면 자동조회 버튼이 활성화됩니다.')
    st.caption('사업보고서·반기보고서 등을 제출하지 않는 법인은 인원 공개 데이터가 없을 수 있습니다. 해당 법인은 수동 입력 또는 확인된 기업목록을 사용하세요.')

    with st.expander('👥 직원 수 엑셀/CSV 일괄 등록 (API 미공개 법인 보완)', expanded=False):
        st.caption('기존 기업목록에 직원 수가 없으면 Excel/CSV로 보완합니다. 반드시 실제 확인된 인원만 입력하세요. ' 
                   '직원 수가 없는 기업의 시행연도는 자동 추정하지 않습니다.')
        example = pd.DataFrame([{'법인명': '예시회사', 'DART고유번호': '', '종업원수': 65}])
        st.download_button('CSV 입력양식', data=example.to_csv(index=False).encode('utf-8-sig'),
                           file_name='퇴직연금_직원수_입력양식.csv', mime='text/csv', key='pension_headcount_template')
        upload = st.file_uploader('직원 수 자료 업로드', type=['xlsx', 'csv'], key='pension_headcount_upload')
        if upload is not None:
            try:
                mapping, missed, total = _pension_target_import_headcount(upload, corp_master)
                st.caption(f'총 {total}행 / DART 고유번호·법인명 매칭 {len(mapping)}개 / 매칭 불가·인원 무효 {missed}행 (중복 법인은 최종 값 적용)')
                if mapping and st.button('✅ 확인된 직원 수 DB에 저장', key='pension_headcount_save'):
                    _pension_target_save_headcount(mapping)
                    if not _db_cloud_push():
                        st.warning('로컬 DB에는 저장했지만 Supabase 동기화가 확인되지 않았습니다. DB 관리에서 영구저장을 재실행하세요.')
                    st.success(f'{len(mapping)}개 법인의 직원 수 저장 완료')
                    st.rerun()
            except Exception as exc:
                st.error(f'직원 수 파일 검증 오류: {exc}')

    f1, f2, f3, f4 = st.columns([1.3, 1, 1, 1.3])
    types = sorted(df['대상유형'].dropna().unique().tolist())
    selected_types = f1.multiselect('분류 필터', types, default=[t for t in types if t in ('미도입 후보', '가입 여부 확인', '자료 부족', '미분석')],
                                    key='pension_target_filter_type')
    years = f2.selectbox('적용 예정연도', ['전체', '2028년', '2029년', '2030년', '인원 확인 필요'], key='pension_target_year')
    stages = f3.selectbox('영업 진행상태', ['전체'] + PENSION_LEAD_STAGES, key='pension_target_stage')
    name_filter = f4.text_input('법인명 / 소재지 검색', key='pension_target_search', placeholder='청주, 제조, 법인명 등')
    view = df.copy()
    if selected_types:
        view = view[view['대상유형'].isin(selected_types)]
    else:
        view = view.iloc[:0]
    if years != '전체':
        view = view[view['시행예정'].str.startswith(years)]
    if stages != '전체':
        view = view[view['영업상태'] == stages]
    if name_filter:
        search = str(name_filter).strip()
        view = view[view['법인명'].str.contains(search, case=False, regex=False, na=False) |
                    view['주소'].str.contains(search, case=False, regex=False, na=False)]
    st.caption(f'현재 필터 결과: {len(view):,}개. 미분석·기존 운영 법인도 보고 싶으면 **분류 필터**에서 선택하세요.')
    shown_cols = ['법인명', '주소', '종업원수', '인원출처', '인원공시기준', '시행예정', '대상유형', '판별근거', '영업상태', '접수일', '분석저장일', '출처']
    st.dataframe(view[shown_cols], use_container_width=True, hide_index=True, height=425)
    st.download_button('📥 현재 조건의 타깃 법인 Excel 다운로드', data=_pension_target_excel(view),
                       file_name='CorporateFinder_퇴직연금_영업타깃.xlsx',
                       mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                       key='pension_targets_excel', disabled=view.empty)

    if not view.empty:
        st.markdown('#### 📋 법인별 공시 근거 및 영업 진행 기록')
        choices = [f'{row["법인명"]} · {row["DART고유번호"] or "매칭 필요"}'
                   for row in view.to_dict('records')]
        picked = st.selectbox('법인 선택', range(len(choices)), format_func=lambda x: choices[x], key='pension_selected')
        r = view.iloc[picked]
        st.write(f'**{r["대상유형"]}** | {r["판별근거"]}')
        st.write(f'**제안 방향:** {r["영업제안"]}')
        if r['DART원문'] and str(r['DART원문']).startswith('https://dart.fss.or.kr/'):
            st.link_button('해당 감사보고서 원문 확인', r['DART원문'])
        code = str(r['DART고유번호'])
        if code:
            with get_db_connection() as conn:
                item = conn.execute('SELECT result_json FROM company_analysis WHERE corp_code=?', (code,)).fetchone()
            if item:
                try:
                    pen = (json.loads(item['result_json']).get('pension') or {})
                    op = pen.get('운영근거') or {}
                    with st.expander('명시적 운영/미운용 판별 원문', expanded=False):
                        for key in ('미운용', 'DB', 'DC'):
                            evid = op.get(key) or []
                            if evid:
                                st.markdown(f'**{key} ({len(evid)}건)**')
                                for line in evid[:4]:
                                    st.info(str(line))
                except (ValueError, TypeError):
                    pass
            with st.form('pension_target_contact_form'):
                count_value = r['종업원수']
                st.caption('직원 수는 수동으로 검증 후 등록 (퇴직연금 도입 여부 판별값과 별도)')
                count_text = st.text_input('확인된 종업원 수 (모르면 공란)', value='' if pd.isna(count_value) else str(int(count_value)))
                current_stage = str(r['영업상태'])
                current_stage = current_stage if current_stage in PENSION_LEAD_STAGES else '미접촉'
                stage = st.selectbox('영업 진행상태', PENSION_LEAD_STAGES,
                                     index=PENSION_LEAD_STAGES.index(current_stage))
                note = st.text_area('상담메모', value=str(r['상담메모']), height=90)
                if st.form_submit_button('💾 영업정보 저장', type='primary'):
                    try:
                        headcount = _pension_int(count_text)
                        if count_text.strip() and headcount is None:
                            raise ValueError('종업원 수는 정확한 정수만 입력해 주세요.')
                        _pension_target_save(code, headcount=headcount, stage=stage, note=note)
                        if not _db_cloud_push():
                            st.warning('로컬 DB 저장 완료. Supabase 동기화는 실패/미설정 상태이므로 DB 관리에서 별도 확인하세요.')
                        st.success('영업 진행상태·직원 수·메모를 저장했습니다.')
                        st.rerun()
                    except Exception as exc:
                        st.error(f'저장 실패: {exc}')
        else:
            st.warning('동명 또는 미매칭 법인입니다. DART 법인을 정확히 선택한 후 저장할 수 있습니다.')
    st.caption('※ 직원 현황은 공시 기준 참고 인원으로 실제 상시근로자 수와 다를 수 있습니다. 이 탭의 분류는 영업 검토 보조용이며 가입 여부, 기존 연금사업자, 퇴직급여채무 금액, 법정 적립률을 확정하지 않습니다.')

# LOGIN — MATCHES HANAV ETF FINDER
# =========================================================
st.markdown("""
<style>
:root {color-scheme:light;--primary-color:#008878;--background-color:#F4F8F7}
.stApp,[data-testid="stAppViewContainer"]{background:#F4F8F7;color:#203D35}
.block-container{max-width:1500px;padding-top:1rem}
input{color:#203D35!important}
.stButton>button[kind="primary"]{background:#008878!important;border-color:#008878!important;color:white!important}
</style>
""",unsafe_allow_html=True)

def _corp_app_password():
    return str(get_secret('APP_PASSWORD','') or '')

def _corp_check_password():
    expected=_corp_app_password()
    entered=str(st.session_state.get('_corp_password',''))
    ok=bool(expected) and hmac.compare_digest(entered.encode('utf-8'),expected.encode('utf-8'))
    st.session_state['_corp_authenticated']=ok
    st.session_state['_corp_login_error']=not ok
    st.session_state['_corp_password']=''

def _corp_logout():
    for key in list(st.session_state):
        del st.session_state[key]

st.session_state.setdefault('_corp_authenticated',False)
st.session_state.setdefault('_corp_login_error',False)
if not _corp_app_password():
    st.error('APP_PASSWORD가 설정되지 않았습니다. Streamlit Secrets에 APP_PASSWORD를 등록해 주세요.')
    st.stop()
if not st.session_state['_corp_authenticated']:
    st.markdown('''
    <div style="max-width:430px;margin:11vh auto 22px;text-align:center;">
       <div style="font-size:34px;font-weight:950;color:#008878;">Corporate Finder</div>
       <div style="margin-top:7px;color:#536F64;font-size:13px;font-weight:650;">Private Corporate Intelligence Dashboard</div>
       <div style="margin-top:4px;color:#60766E;font-size:11px;">Designed &amp; Built by K.H. Ahn</div>
    </div>''',unsafe_allow_html=True)
    _left,_center,_right=st.columns([1,1.15,1])
    with _center:
        with st.form('_corp_login_form'):
            st.text_input('비밀번호',type='password',key='_corp_password',placeholder='비밀번호를 입력하세요')
            st.form_submit_button('🔐 로그인',type='primary',use_container_width=True,on_click=_corp_check_password)
        if st.session_state.get('_corp_login_error'):
            st.error('비밀번호가 올바르지 않습니다.')
    st.stop()

_, _logout_col=st.columns([5,1])
with _logout_col:
    st.button('🔒 로그아웃',key='_corp_logout_btn',use_container_width=True,on_click=_corp_logout)
st.title('🏢 Corporate Finder')
st.caption('법인명·본점 소재지 검색 · 감사보고서 거래은행/퇴직연금 · 마지막 DB 자동 복원')

# =========================================================
# 로컬 DB 먼저 실행: 온라인 업데이트는 명시적으로 요청할 때만
# =========================================================
api_key = str(get_secret("DART_API_KEY", "")).strip()
offline_mode = st.checkbox(
    "오프라인 모드 (저장된 DB만 사용 / DART 접속 안 함)",
    value=not bool(api_key),
    key="corp_finder_offline_mode",
)
dart_enabled = bool(api_key) and not offline_mode

# 로컬 DB가 존재하는 PC에서는 원격 DB 다운로드보다 화면 표시를 우선합니다.
# 새 배포 환경처럼 로컬 법인 DB가 없을 때만 Supabase를 자동 복원합니다.
try:
    if _db_has_usable_local_master():
        _db_startup_state = 'local_ready'
    else:
        with st.spinner('저장된 법인 DB를 확인하고 있습니다. 로컬 DB가 없어 Supabase 백업을 복원하는 중입니다...'):
            _db_startup_state = _db_startup_restore()
    init_local_db()
    corp_df = load_saved_corp_master()
except Exception as e:
    st.error(f"로컬 DB 초기화/불러오기 실패: {type(e).__name__}: {e}")
    st.caption(f"확인 중인 DB 경로: {DB_PATH}")
    st.stop()

if corp_df.empty:
    st.warning('현재 폴더에 조회 가능한 법인 DB가 없습니다. app.py만 다른 폴더로 옮기면 이전 폴더의 DB는 자동으로 복사되지 않습니다.')
    st.caption(f'현재 DB 경로: {DB_PATH}')
    if st.session_state.get('_db_cloud_error'):
        st.caption('Supabase 자동 복원에 실패했습니다. 아래 버튼으로 재시도하거나 기존 DB를 업로드할 수 있습니다.')

    if st.button('☁️ Supabase 저장 DB 다시 가져오기', key='startup_cloud_restore_retry', type='primary'):
        try:
            with st.spinner('Supabase에 저장된 전체 법인 DB를 가져오는 중입니다...'):
                result = _db_cloud_restore_on_boot()  # 캐시를 통하지 않는 명시적 재시도
            if result in ('restored', 'seeded'):
                st.session_state.pop('_db_cloud_error', None)
                st.success('법인 DB를 복구했습니다.')
                st.rerun()
            else:
                st.warning('Supabase에서 복구할 DB를 확인하지 못했습니다. 설정과 네트워크를 점검해 주세요.')
        except Exception as exc:
            st.error(f'클라우드 DB 복구 실패: {type(exc).__name__}')

    if dart_enabled and st.button('🌐 DART 법인 고유번호 목록 최초 구축', key='startup_dart_master', type='secondary'):
        try:
            with st.spinner('DART에서 법인 고유번호 목록을 수집하는 중입니다...'):
                fresh_df = load_corp_codes(api_key)
                sync_corp_master(fresh_df)
                set_db_meta('corp_master_updated_at', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
                _db_cloud_push()
            st.success('법인목록을 구축했습니다.')
            st.rerun()
        except Exception as exc:
            st.error(f'DART 최초 법인목록 구축 실패: {type(exc).__name__}: {exc}')

    startup_restore = st.file_uploader('기존 corp_finder_cache.sqlite3 수동 복원', type=['sqlite3', 'db'], key='startup_full_restore')
    if startup_restore is not None and st.button('업로드한 DB 복원', key='startup_full_restore_btn'):
        try:
            restore_full_db_bytes(startup_restore.getvalue())
            st.success('로컬 DB를 복원했습니다. 화면을 새로고침합니다.')
            st.rerun()
        except Exception as exc:
            st.error(f'복원 실패: {type(exc).__name__}: {exc}')
    st.stop()

_db_cloud_indicator()
if 'external_company_df' not in st.session_state:
    try:
        _saved_external, _saved_columns = load_external_company_db()
        if not _saved_external.empty:
            st.session_state['external_company_df'] = _saved_external
            st.session_state['external_mapped_cols'] = _saved_columns
    except Exception as _ext_exc:
        st.warning(f'저장된 외부 기업목록 불러오기 실패: {_ext_exc}')

last_db_sync = get_db_meta("corp_master_updated_at", "기존 저장본 (갱신일 기록 없음)")
mode_label = "오프라인" if not dart_enabled else "로컬 DB 우선"
st.caption(f"💾 {mode_label} · 저장 법인 {len(corp_df):,}개 · 마지막 법인목록 갱신: {last_db_sync} · DB 위치: {DB_PATH.name}")
if st.button("🔄 DART 법인목록 최신 업데이트", disabled=not dart_enabled, key="refresh_master_db"):
    try:
        with st.spinner("DART에서 최신 법인목록을 수집 중입니다. 저장된 DB는 정상 수집 후 갱신합니다..."):
            load_corp_codes.clear()  # 이전 Streamlit 캐시가 아닌 새 DART 목록 요청
            fresh_df = load_corp_codes(api_key)
            sync_corp_master(fresh_df)
            set_db_meta("corp_master_updated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            _db_cloud_push()
        st.success("최신 법인목록을 로컬 DB에 저장했습니다.")
        st.rerun()
    except Exception as exc:
        st.error(f"업데이트에 실패했습니다. 기존 저장 DB를 유지합니다: {exc}")

if not dart_enabled:
    st.info("저장된 법인·주소·전화번호를 검색할 수 있습니다. 거래은행·퇴직연금은 저장된 분석이 있는 법인에 한해 조회 가능합니다. PC에서 새로 분석했다면 상단 'PC 최신 분석 결과 가져오기'를 누르세요.")


# =========================================================
# UI
# =========================================================
name_tab, region_tab, pension_target_tab, bulk_tab, db_tab = st.tabs([
    "🔎 법인명 조회", "📍 본점 소재지 검색", "🎯 퇴직연금 타깃", "📚 여러 법인 일괄 조회", "🗄️ 주소 DB 관리"
])


with name_tab:
    st.subheader("회사명으로 검색")
    query = st.text_input(
        "법인명",
        placeholder="예: 다셀주식회사, 우진교통, 신흥기업",
        key="single_query",
    )

    if query:
        matches = search_companies(corp_df, query)
        if matches.empty:
            st.warning("DART 고유번호 목록에서 회사를 찾지 못했습니다.")
        else:
            options = matches.to_dict("records")
            labels = [company_label(x) for x in options]
            selected_label = st.selectbox("회사 선택", labels, key="name_company_select")
            selected = options[labels.index(selected_label)]

            info = get_saved_company_profile(selected["corp_code"])
            if info is None and dart_enabled:
                try:
                    info = get_company_info(api_key, selected["corp_code"])
                    save_company_profile(selected, info, "")
                except Exception:
                    pass

            if info:
                cols = st.columns(4)
                cols[0].metric("대표자", info.get("ceo_nm") or "-")
                cols[1].metric("사업자번호", info.get("bizr_no") or "-")
                cols[2].metric("법인번호", info.get("jurir_no") or "-")
                cols[3].metric("업종코드", info.get("induty_code") or "-")
                if info.get("adres"):
                    st.write(f"**본점 주소:** {info['adres']}")
                st.write(f"**전화번호:** {info.get('phn_no') or '전화번호 미공개'}")

            _name_c1, _name_c2 = st.columns([2, 1])
            with _name_c1:
                _name_run = st.button("거래은행 + 퇴직연금 통합 분석", type="primary", key="single_analyze")
            with _name_c2:
                _name_refresh = st.button("🔄 최신 공시 재분석", key="single_analyze_refresh", disabled=not dart_enabled)
            if _name_run or _name_refresh:
                try:
                    with st.spinner("저장 결과 확인 또는 DART 감사보고서 분석 중입니다..."):
                        result = analyze_company(api_key, selected, refresh=_name_refresh)
                    st.session_state["single_result"] = result
                    if dart_enabled and result.get('report') and not _pension_employee_saved(selected['corp_code']):
                        _prior=_pension_employee_recent_checks().get(str(selected['corp_code']))
                        _recent=bool(_prior and str(_prior.get('checked_at','')) >=
                            (datetime.now()-timedelta(days=30 if _prior.get('status')=='미공개' else 1)).isoformat(timespec='seconds'))
                        if not _recent:
                            try:
                                with st.spinner('DART 정기보고서 직원현황 추가 조회 중...'):
                                    _record=_pension_employee_lookup(api_key,selected['corp_code'])
                                if _record:
                                    _pension_employee_save(selected['corp_code'],_record)
                                    _db_cloud_push()
                                else:
                                    st.caption('직원 현황: DART 정기보고서 데이터 미공개. 타깃 탭에서 수동 확인 가능')
                            except Exception as _err:
                                _pension_employee_log(selected_corp['corp_code'],'조회 오류',type(_err).__name__)
                                st.caption(f'직원 수 추가 조회 불가 ({type(_err).__name__}), 감사보고서 분석은 정상 표시됩니다.')
                except Exception as e:
                    st.error(f"분석 중 오류: {e}")

    result = st.session_state.get("single_result")
    if result:
        st.divider()

        c = result["company"]
        r = result.get("report")

        st.subheader(c["법인명"])
        info_cols = st.columns(4)
        info_cols[0].metric("대표자", c.get("대표자") or "-")
        info_cols[1].metric("사업자번호", c.get("사업자등록번호") or "-")
        info_cols[2].metric("DART 코드", c.get("DART고유번호") or "-")
        info_cols[3].metric("확인 은행 수", len(result.get("banks", [])))

        if c.get("주소"):
            st.write(f"**주소:** {c['주소']}")
        st.write(f"**전화번호:** {c.get('전화번호') or '전화번호 미공개'}")

        if r:
            st.write(f"**최근 감사보고서:** {r['보고서명']} · 접수일 {r['접수일']}")
            st.link_button("DART 원문 열기", r["DART링크"])
        else:
            st.warning(result.get("error", "감사보고서를 찾지 못했습니다."))

        banks = result.get("banks", [])
        if banks:
            render_bank_analysis(banks, "🏦 감사보고서에서 확인된 금융기관")
        elif r:
            st.warning("감사보고서 원문에서 등록된 은행명을 찾지 못했습니다.")

        if r:
            _render_pension_employee_section(str(c.get('DART고유번호') or ''),api_key,dart_enabled,key='single_pension')
            render_pension_analysis(result.get("pension") or analyze_pension_text(""))

        generic = result.get("generic", [])
        if generic:
            with st.expander("은행명이 특정되지 않은 금융기관 관련 문구"):
                for txt in generic[:10]:
                    st.write("• " + txt)

        render_visit_report_exports(result, 'single_visit')

        excel = make_excel([result])
        st.download_button(
            "📥 기존 상세 분석 엑셀 다운로드",
            data=excel,
            file_name=f"{c['법인명']}_법인Finder.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


with region_tab:
    st.subheader("본점 소재지 / 사업장 소재지로 검색")
    st.caption(
        "DART 주소 DB 또는 외부 기업목록(CSV/Excel)에서 회사를 찾은 뒤, "
        "DART 법인과 매칭하여 거래은행·차입금·담보·예금·퇴직연금 공시를 분석합니다."
    )

    total_cnt, addr_cnt, fail_cnt = address_db_stats()
    c1, c2, c3 = st.columns(3)
    c1.metric("DART 전체 법인", f"{total_cnt:,}")
    c2.metric("주소 DB 구축", f"{addr_cnt:,}")
    c3.metric("수집 오류", f"{fail_cnt:,}")

    if addr_cnt == 0:
        st.warning(
            "현재 DART 주소 DB가 0건입니다. DART 전체 법인목록에는 주소가 포함되어 있지 않아 "
            "본점 소재지 검색을 하려면 회사별 기업개황을 먼저 수집해야 합니다."
        )
        qa, qb = st.columns([1.2, 2.8])
        quick_batch = qa.selectbox(
            "빠른 주소수집", [50, 100, 200, 300], index=1, key="region_quick_batch"
        )
        if qb.button("🚀 DART 주소 DB 지금 구축", type="primary", key="region_quick_build", disabled=not dart_enabled):
            ok, fail = build_address_cache(api_key, quick_batch, True, False)
            if ok or fail:
                st.success(f"주소 DB 구축: 성공 {ok:,}개 / 오류 {fail:,}개")
                st.rerun()
            else:
                st.info("추가로 수집할 법인이 없습니다.")
        st.info(
            "청주시 제조업체를 빠르게 찾는 목적이라면 DART 주소 DB를 11만 건 넘게 구축하는 것보다 "
            "아래 '외부 기업목록'에 충청북도/청주시 제조업체 총람을 올리는 방식이 훨씬 빠릅니다. "
            "외부 목록에서 회사를 고른 뒤 DART 법인과 매칭해 감사보고서·거래은행 분석으로 연결할 수 있습니다."
        )
    elif addr_cnt < 1000:
        st.caption(
            f"현재 본점주소 {addr_cnt:,}건만 수집되어 있어 특정 지역의 결과가 일부 누락될 수 있습니다. "
            "더 넓게 찾으려면 주소 DB를 추가 구축하거나 외부 기업목록을 함께 사용하세요."
        )

    # 처음 실행할 때는 별도 파일 업로드를 요구하지 않는 DART 주소 DB를 기본값으로 둡니다.
    # 외부 목록은 사용자가 명시적으로 선택했을 때만 업로드 영역을 보여 줍니다.
    source_mode = st.radio(
        "검색 데이터",
        ["DART 주소 DB", "외부 기업목록", "둘 다"],
        index=2 if isinstance(st.session_state.get('external_company_df'), pd.DataFrame)
                 and not st.session_state['external_company_df'].empty else 0,
        horizontal=True,
        key="region_source_mode",
        help="DART 주소 DB=본점 소재지, 외부 기업목록=본점/공장/사업장 등 파일에 들어 있는 주소 기준",
    )

    if source_mode == "DART 주소 DB" and addr_cnt == 0:
        st.info(
            "현재는 DART 본점주소가 아직 0건입니다. 위의 '🚀 DART 주소 DB 지금 구축'을 먼저 눌러 주세요. "
            "주소는 DART 전체 법인목록에 포함되어 있지 않아 회사별 기업개황을 한 번씩 수집해야 합니다."
        )

    if source_mode in ("외부 기업목록", "둘 다"):
        uploaded_files = st.file_uploader(
            "기업목록 CSV / Excel (여러 파일 추가 가능)",
            type=["csv", "xlsx", "xls"],
            accept_multiple_files=True,
            key="external_company_files_multi",
            help="여러 파일을 한 번에 선택하거나 + 버튼으로 추가할 수 있습니다. 기존 DB 목록은 삭제되지 않습니다.",
        )
        st.caption("새 파일을 추가하면 기존 목록에 누적 저장합니다. 법인명과 사업장 주소가 같은 행은 중복 추가하지 않습니다.")

        if uploaded_files:
            already_processed = set(st.session_state.get("external_uploaded_hashes", []))
            old_df = st.session_state.get("external_company_df")
            merged_df = old_df.copy() if isinstance(old_df, pd.DataFrame) else pd.DataFrame()
            saved_mapping = dict(st.session_state.get("external_mapped_cols", {}))
            stored_info = saved_mapping.get("__파일인식정보__", [])
            history = ([stored_info] if isinstance(stored_info, dict)
                       else list(stored_info) if isinstance(stored_info, list) else [])
            added_total, skipped_total = 0, 0
            completed_hashes = []

            for external_upload in uploaded_files:
                file_hash = hashlib.sha256(external_upload.getvalue()).hexdigest()
                if file_hash in already_processed or file_hash in completed_hashes:
                    continue
                try:
                    imported_df, mapped = read_external_company_file(external_upload)
                    previous_count = len(merged_df)
                    merged_df = merge_external_company_rows(merged_df, imported_df)
                    added = len(merged_df) - previous_count
                    added_total += added
                    skipped_total += len(imported_df) - added
                    if added:
                        saved_mapping.update({k: v for k, v in mapped.items() if k != "__파일인식정보__"})
                        info = dict(mapped.get("__파일인식정보__", {}))
                        info["추가건수"] = added
                        history.append(info)
                    completed_hashes.append(file_hash)
                except Exception as e:
                    st.error(f"'{external_upload.name}' 읽기 실패 (기존 자료 유지): {e}")

            if completed_hashes:
                if added_total:
                    saved_mapping["__파일인식정보__"] = history
                    try:
                        # 누적 결합본을 한 번만 저장하여 부분 덮어쓰기를 방지한다.
                        save_external_company_db(merged_df, saved_mapping)
                    except Exception as e:
                        st.error(f"외부 기업목록 누적 저장 실패 (기존 저장본 유지): {e}")
                    else:
                        st.session_state["external_company_df"] = merged_df
                        st.session_state["external_mapped_cols"] = saved_mapping
                        st.session_state["external_uploaded_hashes"] = list(
                            already_processed.union(completed_hashes)
                        )
                        st.success(
                            f"신규 {added_total:,}건 추가 · 중복 {skipped_total:,}건 제외 "
                            f"· 누적 {len(merged_df):,}건 SQLite 저장 완료"
                        )
                        if st.session_state.get("_db_cloud_error"):
                            st.warning("Supabase 클라우드 백업이 실패했습니다. '현재 DB 영구 저장'에서 재시도해 주세요.")
                else:
                    st.session_state["external_uploaded_hashes"] = list(
                        already_processed.union(completed_hashes)
                    )
                    st.info(f"중복 {skipped_total:,}건 확인 · 기존 저장 목록을 그대로 유지했습니다.")

        saved_ext = st.session_state.get("external_company_df")
        if isinstance(saved_ext, pd.DataFrame) and not saved_ext.empty:
            st.success(f"💾 저장된 외부 기업목록 누적 {len(saved_ext):,}건 · 파일 재업로드 불필요")
            with st.expander("업로드 파일 및 인식된 열 확인"):
                st.json(st.session_state.get("external_mapped_cols", {}))

        st.link_button(
            "📥 2025 충청북도 제조업체 총람 다운로드 페이지",
            "https://www.chungbuk.go.kr/www/contents.do?key=5382",
        )
        st.caption(
            "사용 순서: 위 페이지에서 '청주시 제조업체총람' 다운로드 → 이 화면의 기업목록에 업로드 → "
            "청주시/구/읍·면·동 검색. 이 자료는 사업장/공장 소재지 기준이며 DART 본점주소와는 의미가 다를 수 있습니다."
        )

    st.markdown("#### 지역 조건")
    r1, r2, r3, r4 = st.columns([1.2, 1, 1, 1])
    province = r1.selectbox("시·도", REGION_PROVINCES, index=REGION_PROVINCES.index("충청북도"), key="region_province")
    city = r2.text_input("시·군", value="청주시" if province == "충청북도" else "", placeholder="예: 청주시", key="region_city")
    if province == "충청북도" and "청주시" in city:
        district = r3.selectbox("구", CHEONGJU_DISTRICTS, key="region_district")
    else:
        district = r3.text_input("구", placeholder="예: 흥덕구", key="region_district_text")
    town = r4.text_input("읍·면·동 / 상세지역", placeholder="예: 오창읍", key="region_town")

    name_filter = st.text_input("법인명 추가 필터 (선택)", placeholder="예: 테크, 산업, 전자", key="region_name_filter")
    max_rows = st.select_slider("최대 검색 결과", options=[100, 300, 500, 1000, 3000], value=500, key="region_max_rows")

    if st.button("🔍 지역 법인 검색", type="primary", key="region_search_btn"):
        terms = region_terms_from_inputs(province, city, district, town)
        result_frames = []

        if source_mode in ("DART 주소 DB", "둘 다"):
            if addr_cnt == 0:
                st.warning("DART 본점주소 DB가 0건이라 DART 지역검색 결과는 없습니다. 위의 주소 DB 구축을 먼저 실행해 주세요.")
            dart_loc = get_profile_rows_for_search(terms, limit=max_rows)
            if not dart_loc.empty:
                if name_filter:
                    q = normalize_company_name(name_filter)
                    dart_loc = dart_loc[dart_loc["corp_name"].map(normalize_company_name).str.contains(re.escape(q), na=False)]
                dart_show = pd.DataFrame({
                    "출처": "DART",
                    "법인명": dart_loc["corp_name"],
                    "대표자": dart_loc["ceo_nm"],
                    "주소": dart_loc["adres"],
                    "업종": dart_loc["induty_code"],
                    "사업자등록번호": dart_loc["bizr_no"],
                    "법인등록번호": dart_loc["jurir_no"],
                    "DART고유번호": dart_loc["corp_code"],
                    "전화번호": dart_loc["phn_no"],
                })
                result_frames.append(dart_show)

        if source_mode in ("외부 기업목록", "둘 다"):
            ext_df = st.session_state.get("external_company_df")
            if ext_df is None or ext_df.empty:
                if source_mode == "외부 기업목록":
                    st.warning("외부 기업목록을 먼저 업로드해 주세요.")
                else:
                    st.caption("외부 기업목록이 없어 이번 검색에서는 DART 주소 DB만 사용합니다.")
            else:
                ext_loc = filter_external_by_region(ext_df, terms, name_filter, limit=max_rows)
                if not ext_loc.empty:
                    ext_show = pd.DataFrame({
                        "출처": "외부목록",
                        "법인명": ext_loc["법인명"],
                        "대표자": ext_loc["대표자"],
                        "주소": ext_loc["주소"],
                        "업종": ext_loc["업종"],
                        "사업자등록번호": ext_loc["사업자등록번호"],
                        "법인등록번호": ext_loc["법인등록번호"],
                        "DART고유번호": "",
                        "전화번호": ext_loc["전화번호"],
                    })
                    result_frames.append(ext_show)

        if result_frames:
            combined = pd.concat(result_frames, ignore_index=True)
            combined["_dedupe"] = combined["법인명"].map(normalize_company_name) + "|" + combined["주소"].astype(str)
            combined = combined.drop_duplicates("_dedupe").drop(columns="_dedupe").reset_index(drop=True)
            st.session_state["region_search_results"] = combined.head(max_rows)
        else:
            st.session_state["region_search_results"] = pd.DataFrame()

    region_results = st.session_state.get("region_search_results", pd.DataFrame())
    if isinstance(region_results, pd.DataFrame) and not region_results.empty:
        st.markdown(f"### 검색 결과 · {len(region_results):,}개")
        st.caption("표에서 회사를 한 줄 클릭하면 아래에서 DART 법인을 매칭한 뒤 거래은행·차입금 분석을 할 수 있습니다.")
        event = st.dataframe(
            region_results,
            use_container_width=True,
            hide_index=True,
            on_select="rerun",
            selection_mode="single-row",
            key="region_results_grid",
        )

        selected_indices = []
        try:
            selected_indices = list(event.selection.rows)
        except Exception:
            selected_indices = []

        if selected_indices:
            idx = selected_indices[0]
            picked = region_results.iloc[idx].to_dict()
            previous = st.session_state.get("region_picked_company") or {}
            previous_key = (str(previous.get("법인명", "")), str(previous.get("주소", "")))
            picked_key = (str(picked.get("법인명", "")), str(picked.get("주소", "")))
            if picked_key != previous_key:
                # 다른 회사를 클릭했는데 직전 회사의 은행 분석 결과가 남아 보이지 않도록 초기화
                st.session_state.pop("region_analysis_result", None)
            st.session_state["region_picked_company"] = picked

        picked = st.session_state.get("region_picked_company")
        if picked:
            st.markdown("#### 선택 회사")
            st.write(f"**{picked.get('법인명','')}**")
            st.write(f"주소: {picked.get('주소') or '-'}")
            st.write(f"전화번호: {picked.get('전화번호') or '전화번호 미공개'}")
            if picked.get("업종"):
                st.write(f"업종/제품: {picked.get('업종')}")

            direct_code = str(picked.get("DART고유번호", "")).strip()
            candidate_rows = []
            if direct_code:
                hit = corp_df[corp_df["corp_code"].astype(str) == direct_code]
                candidate_rows = hit.to_dict("records")
            else:
                candidate_rows = dart_match_for_name(corp_df, picked.get("법인명", ""))

            if not candidate_rows:
                st.warning("DART에서 같은 이름의 공시대상 법인을 찾지 못했습니다. 외부 목록의 모든 회사가 DART 공시대상인 것은 아닙니다.")
            else:
                labels = [company_label(x) for x in candidate_rows]
                sel_label = st.selectbox("DART 법인 매칭", labels, key="region_dart_match_select")
                selected_corp = candidate_rows[labels.index(sel_label)]
                _selected_code = str(selected_corp.get("corp_code", ""))
                if st.session_state.get("_last_region_dart_code") != _selected_code:
                    st.session_state.pop("region_analysis_result", None)
                    st.session_state["_last_region_dart_code"] = _selected_code

                st.info("이 회사를 선택하면 감사보고서에서 거래은행·금융거래 근거와 퇴직연금 관련 공시를 함께 분석합니다.")
                _rg_c1, _rg_c2, _rg_c3 = st.columns([2.2, 1.2, 1.1])
                with _rg_c1:
                    _rg_run = st.button("🏦 거래은행 · 차입금 · 퇴직연금 분석", type="primary", key="region_analyze_btn")
                with _rg_c2:
                    _rg_refresh = st.button("🔄 최신 공시 재분석", key="region_analyze_refresh", disabled=not dart_enabled)
                with _rg_c3:
                    _rg_diagnose = st.button("🔌 DART 연결 진단", key="region_dart_diagnose", disabled=not dart_enabled)
                if _rg_run or _rg_refresh:
                    try:
                        with st.spinner("저장 결과 확인 또는 DART 감사보고서 분석 중입니다..."):
                            result = analyze_company(api_key, selected_corp, refresh=_rg_refresh)
                        st.session_state["region_analysis_result"] = result
                        if dart_enabled and result.get('report') and not _pension_employee_saved(selected_corp['corp_code']):
                            _prior=_pension_employee_recent_checks().get(str(selected_corp['corp_code']))
                            _recent=bool(_prior and str(_prior.get('checked_at','')) >=
                                (datetime.now()-timedelta(days=30 if _prior.get('status')=='미공개' else 1)).isoformat(timespec='seconds'))
                            if not _recent:
                                try:
                                    with st.spinner('DART 정기보고서 직원현황 추가 조회 중...'):
                                        _record=_pension_employee_lookup(api_key,selected_corp['corp_code'])
                                    if _record:
                                        _pension_employee_save(selected_corp['corp_code'],_record)
                                        _db_cloud_push()
                                    else:
                                        st.caption('직원 현황: DART 정기보고서 데이터 미공개. 타깃 탭에서 수동 확인 가능')
                                except Exception as _err:
                                    st.caption(f'직원 수 추가 조회 불가 ({type(_err).__name__}), 감사보고서 분석은 정상 표시됩니다.')
                    except Exception as e:
                        st.error(f"분석 중 오류: {e}")
                if _rg_diagnose:
                    with st.spinner("DART 국문·영문 공시목록 연결을 각각 확인 중입니다..."):
                        _rg_checks = check_dart_connection(api_key, selected_corp["corp_code"])
                    st.dataframe(pd.DataFrame(_rg_checks), hide_index=True, use_container_width=True)
                    st.caption(
                        "DNS=서버 주소 검색, TCP 443=네트워크 접속, TLS=HTTPS 보안 협상, "
                        "공시목록 API=실제 DART 응답입니다. DART 000=조회 성공, 013=자료 없음, "
                        "012=접속 IP 제한, 020=요청 한도 초과. "
                        "TCP/TLS는 직접 연결 테스트이므로 네트워크 프록시를 사용하는 환경에서는 "
                        "API 결과와 달라질 수 있습니다."
                    )
                    _rg_ok = any(
                        item["단계"] == "공시목록 API" and (
                            "DART 000" in item["결과"] or "DART 013" in item["결과"]
                        ) for item in _rg_checks
                    )
                    if not _rg_ok:
                        st.warning(
                            "DART 실제 API 연결이 확인되지 않습니다. 두 서버 모두 TCP 단계에서 실패한다면 "
                            "Streamlit Cloud의 외부 연결 제한이나 DART 측 접속 IP 정책을 점검해야 합니다. "
                            "기존 Supabase 저장 분석은 계속 조회할 수 있습니다."
                        )

        rr = st.session_state.get("region_analysis_result")
        if rr:
            st.divider()
            c = rr["company"]
            r = rr.get("report")
            st.markdown(f"### {c['법인명']} 분석")
            cols = st.columns(4)
            cols[0].metric("대표자", c.get("대표자") or "-")
            cols[1].metric("사업자번호", c.get("사업자등록번호") or "-")
            cols[2].metric("법인번호", c.get("법인등록번호") or "-")
            cols[3].metric("확인은행", len(rr.get("banks", [])))
            st.write(f"**본점 주소:** {c.get('주소') or '-'}")
            st.write(f"**전화번호:** {c.get('전화번호') or (picked.get('전화번호') if picked else '') or '전화번호 미공개'}")
            if r:
                st.write(f"**최근 감사보고서:** {r['보고서명']} · {r['접수일']}")
                st.link_button("DART 원문 열기", r["DART링크"], key="region_dart_link")
            if rr.get("banks"):
                render_bank_analysis(rr["banks"], "🏦 감사보고서에서 확인된 거래 금융기관")
            elif r:
                st.warning("감사보고서 원문에서 등록된 은행명을 찾지 못했습니다.")
            elif rr.get("error"):
                st.warning(rr["error"])

            if r:
                _render_pension_employee_section(str(c.get('DART고유번호') or ''),api_key,dart_enabled,
                    external_count=(picked.get('종업원수') if picked else None),key='region_pension')
                render_pension_analysis(rr.get("pension") or analyze_pension_text(""))

            generic = rr.get("generic", [])
            if generic:
                with st.expander("은행명이 특정되지 않은 금융기관 관련 문구"):
                    for txt in generic[:10]:
                        st.write("• " + txt)

            render_visit_report_exports(rr, 'region_visit', external_meta=picked)

            st.download_button(
                "📥 기존 상세 분석 엑셀",
                data=make_excel([rr]),
                file_name=f"{c['법인명']}_법인Finder.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="region_excel_download",
            )

        st.caption("지역 검색은 기업을 찾는 단계이고, 실제 거래은행 정보는 선택한 법인의 최신 감사보고서를 분석한 뒤 표시됩니다.")

        csv_data = region_results.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")
        st.download_button(
            "📥 지역 검색결과 CSV 다운로드",
            data=csv_data,
            file_name="법인Finder_지역검색결과.csv",
            mime="text/csv",
            key="region_csv_download",
        )
    elif "region_search_results" in st.session_state:
        ext_ready = isinstance(st.session_state.get("external_company_df"), pd.DataFrame) and not st.session_state.get("external_company_df").empty
        if source_mode == "DART 주소 DB" and addr_cnt == 0:
            st.info(
                "아직 검색할 DART 본점주소 데이터가 없습니다. 화면 위쪽의 '🚀 DART 주소 DB 지금 구축'을 먼저 실행해 주세요."
            )
        elif source_mode == "외부 기업목록" and not ext_ready:
            st.info("먼저 청주시 제조업체총람 같은 CSV/Excel 기업목록을 업로드해 주세요.")
        elif source_mode == "둘 다" and addr_cnt == 0 and not ext_ready:
            st.info(
                "현재 검색 가능한 주소 데이터가 없습니다. DART 주소 DB를 구축하거나 외부 기업목록을 업로드하면 지역검색을 시작할 수 있습니다."
            )
        else:
            st.info(
                "조건에 맞는 회사를 찾지 못했습니다. DART 주소 DB가 일부만 구축된 경우 결과가 누락될 수 있으므로 "
                "주소 DB를 더 구축하거나 외부 기업목록을 함께 사용해 주세요."
            )


with pension_target_tab:
    render_pension_target_tab(corp_df)


with bulk_tab:
    st.subheader("여러 회사 일괄 분석")
    bulk_text = st.text_area(
        "회사명을 한 줄에 하나씩 입력",
        placeholder="다셀주식회사\n우진교통\n신흥기업",
        height=170,
    )
    st.caption("한 번에 최대 20개 법인을 권장합니다.")

    if st.button("일괄 분석 시작", type="primary", key="bulk_analyze"):
        names = [x.strip() for x in bulk_text.splitlines() if x.strip()]
        names = list(dict.fromkeys(names))[:20]

        if not names:
            st.warning("회사명을 입력해 주세요.")
        else:
            results = []
            progress = st.progress(0)
            status = st.empty()

            for idx, name in enumerate(names, start=1):
                status.write(f"{idx}/{len(names)} · {name} 확인 중...")
                matches = search_companies(corp_df, name)

                if matches.empty:
                    results.append({
                        "company": {
                            "법인명": name,
                            "DART고유번호": "",
                            "대표자": "",
                            "사업자등록번호": "",
                            "법인등록번호": "",
                            "주소": "",
                            "홈페이지": "",
                        },
                        "report": None,
                        "banks": [],
                        "generic": [],
                        "error": "DART 법인 매칭 실패",
                    })
                elif len(matches) > 1:
                    results.append({
                        "company": {
                            "법인명": name,
                            "DART고유번호": "",
                            "대표자": "",
                            "사업자등록번호": "",
                            "법인등록번호": "",
                            "주소": "",
                            "홈페이지": "",
                        },
                        "report": None,
                        "banks": [],
                        "generic": [],
                        "error": f"동명/유사 법인 {len(matches)}개 - 법인명 조회에서 정확한 법인을 선택해 주세요.",
                    })
                else:
                    selected = matches.iloc[0].to_dict()
                    try:
                        result = analyze_company(api_key, selected)
                        results.append(result)
                        # 감사보고서 분석 이후 company.json을 중복 호출하지 않는다.
                    except Exception as e:
                        results.append({
                            "company": {
                                "법인명": selected.get("corp_name", name),
                                "DART고유번호": selected.get("corp_code", ""),
                                "대표자": "",
                                "사업자등록번호": "",
                                "법인등록번호": "",
                                "주소": "",
                                "홈페이지": "",
                            },
                            "report": None,
                            "banks": [],
                            "generic": [],
                            "error": str(e),
                        })

                progress.progress(idx / len(names))

            status.empty()
            st.session_state["bulk_results"] = results

    results = st.session_state.get("bulk_results", [])
    if results:
        st.divider()
        st.markdown("### 📊 일괄 분석 결과")

        rows = []
        for result in results:
            c = result["company"]
            r = result.get("report") or {}
            banks = result.get("banks", [])
            rows.append({
                "법인명": c.get("법인명", ""),
                "대표자": c.get("대표자", ""),
                "주소": c.get("주소", ""),
                "최근감사보고서": r.get("보고서명", ""),
                "접수일": r.get("접수일", ""),
                "확인은행": ", ".join(b["은행"] for b in banks),
                "확인은행수": len(banks),
                "퇴직연금":(result.get("pension") or {}).get("운영분류", "확인 불가"),
                "비고": result.get("error", ""),
            })

        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)

        excel = make_excel(results)
        st.download_button(
            "📥 전체 결과 엑셀 다운로드",
            data=excel,
            file_name="법인Finder_일괄분석.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


with db_tab:
    st.subheader("DART 본점주소 DB 구축 / 백업")
    st.markdown("#### 💾 전체 DB 백업 / 복원 (사내망 PC 이동용)")
    st.caption("법인목록·주소·전화번호·기존 거래은행/퇴직연금 분석 결과를 한 파일로 보관합니다. DB 복원 시 현재 데이터를 대체합니다.")
    st.download_button(
        "📥 전체 DB(SQLite) 백업 다운로드",
        data=export_full_db_bytes(),
        file_name="corp_finder_cache.sqlite3",
        mime="application/x-sqlite3",
        key="full_db_export",
    )
    full_restore_file = st.file_uploader("전체 DB 백업파일 복원", type=["sqlite3", "db"], key="full_db_restore_file")
    if full_restore_file is not None and st.button("전체 DB 복원 실행", key="full_db_restore_btn"):
        try:
            restore_full_db_bytes(full_restore_file.getvalue())
            _db_cloud_push()
            st.success("전체 DB 복원 완료. 법인 목록을 다시 불러옵니다.")
            st.rerun()
        except Exception as exc:
            st.error(f"전체 DB 복원 실패: {exc}")
    st.divider()
    total_cnt, addr_cnt, fail_cnt = address_db_stats()
    cols = st.columns(4)
    cols[0].metric("DART 법인", f"{total_cnt:,}")
    cols[1].metric("주소 수집 완료", f"{addr_cnt:,}")
    cols[2].metric("수집 오류", f"{fail_cnt:,}")
    pct = (addr_cnt / total_cnt * 100) if total_cnt else 0
    cols[3].metric("진행률", f"{pct:.1f}%")

    st.info(
        "DART 고유번호 목록에는 주소가 없어서 기업개황 API를 회사별로 조회해야 합니다. "
        "따라서 전체 주소 DB는 한 번에 만들기보다 아래 배치 단위로 누적하는 방식입니다. "
        "OpenDART 호출량 제한에 걸릴 수 있으므로 배치 구축이 안전합니다. "
        "법인명 조회/감사보고서 분석을 한 회사는 자동으로 주소 DB에도 저장됩니다."
    )

    b1, b2, b3 = st.columns([1, 1, 1])
    batch_size = b1.selectbox("이번에 수집할 회사 수", [20, 50, 100, 200, 300, 500, 1000], index=2, key="db_batch_size")
    listed_first = b2.checkbox("상장사 우선", value=True, key="db_listed_first")
    retry_errors = b3.checkbox("오류 회사 재시도 포함", value=False, key="db_retry_errors")

    if st.button("주소 DB 이어서 구축", type="primary", key="db_build_btn", disabled=not dart_enabled):
        ok, fail = build_address_cache(api_key, batch_size, listed_first, retry_errors)
        if ok == 0 and fail == 0:
            st.success("추가로 수집할 회사가 없습니다.")
        else:
            st.success(f"이번 실행: 성공 {ok:,}개 / 오류 {fail:,}개")
            st.rerun()

    st.markdown("#### 주소 DB 백업 / 복원")
    st.caption("엑셀은 주소 정보만 백업합니다. 사내망 PC 이동에는 위의 전체 DB(SQLite) 백업을 사용하세요. Streamlit Cloud 로컬 디스크는 영구 보관을 보장하지 않습니다.")

    st.download_button(
        "📥 주소 DB 엑셀 백업",
        data=export_address_cache_excel(),
        file_name="법인Finder_주소DB.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="db_export_btn",
    )

    restore_file = st.file_uploader("주소 DB 백업파일 복원", type=["xlsx", "xls", "csv"], key="db_restore_file")
    if restore_file is not None and st.button("주소 DB 복원 실행", key="db_restore_btn"):
        try:
            count = import_address_cache_file(restore_file)
            _db_cloud_push()
            st.success(f"주소 DB {count:,}건을 복원했습니다.")
            st.rerun()
        except Exception as e:
            st.error(f"복원 실패: {e}")


st.divider()
st.caption(
    "법인 Finder는 DART 및 사용자가 불러온 공공/업무용 기업목록을 결합해 법인 탐색과 DART 공시 분석을 돕는 업무보조 도구입니다. "
    "특정 은행명이 기재되어 있다는 사실만으로 해당 은행을 주거래은행이라고 단정할 수 없으며, "
    "최종 판단은 감사보고서 원문과 회사 확인을 통해 검증해야 합니다."
)
